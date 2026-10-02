"""Explicit outcomes for the local Linux profile; never turn nonfinishes into times."""

from __future__ import annotations

import math
import os
import shutil
from datetime import datetime, timezone
from hashlib import sha256
from pathlib import Path

from openpyxl import load_workbook


def collect_results(round_object, *, input_fn=input):
    if round_object.get("official_outcomes"):
        return True
    rows = round_object["handicap_results"]
    outcomes = {}
    print("\nEnter raw seconds, DNF, DQ, DNS, VOID, or PENALTY <raw seconds> <penalty seconds>.")
    print("Raw seconds run from the competitor's start to block severed. C cancels entry.")
    for row in rows:
        while True:
            value = input_fn(f"  {row['name']}: ").strip().lower()
            if value == "c":
                return False
            try:
                parts = value.split()
                if value in {"dnf", "dq", "dns", "void"}:
                    status, raw, penalty = value, None, None
                elif len(parts) == 3 and parts[0] == "penalty":
                    status, raw, penalty = "penalty", float(parts[1]), float(parts[2])
                else:
                    status, raw, penalty = "completion", float(value), None
                if any(
                    not math.isfinite(number) or number <= 0 or round(number * 1000) < 1
                    for number in (raw, penalty)
                    if number is not None
                ):
                    raise ValueError("time must be finite and positive")
                outcomes[row["local_competitor_id"]] = {
                    "status": status,
                    "raw_time_ms": None if raw is None else round(raw * 1000),
                    "penalty_ms": None if penalty is None else round(penalty * 1000),
                }
                break
            except ValueError:
                print(
                    "Use positive raw seconds or an explicit outcome; penalty requires raw seconds and penalty seconds."
                )
    finishes = {
        row["local_competitor_id"]: outcomes[row["local_competitor_id"]]["raw_time_ms"]
        + row["mark"] * 1000
        + (outcomes[row["local_competitor_id"]]["penalty_ms"] or 0)
        for row in rows
        if outcomes[row["local_competitor_id"]]["status"] in {"completion", "penalty"}
    }
    proposed = {name: 1 + sum(other < clock for other in finishes.values()) for name, clock in finishes.items()}
    print("Suggested finish-clock placings (ties remain tied):", proposed)
    use_proposed = input_fn("Authorize these as the official judge placings? (y/n): ").strip().lower() == "y"
    placings = {}
    names = {row["local_competitor_id"]: row["name"] for row in rows}
    for name in finishes:
        if use_proposed:
            placings[name] = proposed[name]
        else:
            while True:
                try:
                    value = int(
                        input_fn(f"Official placing for {names[name]} [{name}] (equal positions preserve a tie): ")
                    )
                    if value < 1:
                        raise ValueError("positive official placing required")
                    placings[name] = value
                    break
                except ValueError:
                    print("Enter the judge's positive finish position.")
    for name, outcome in outcomes.items():
        outcome["official_placing"] = placings.get(name)
    if input_fn("Confirm these official outcomes and judge placings? (y/n): ").strip().lower() != "y":
        return False
    round_object["official_outcomes"] = outcomes
    refresh_result_projections(round_object)

    return True


def refresh_result_projections(round_object):
    rows = round_object["handicap_results"]
    outcomes = round_object["official_outcomes"]
    round_object["actual_results"] = {
        row["name"]: outcomes[row["local_competitor_id"]]["raw_time_ms"] / 1000
        for row in rows
        if outcomes[row["local_competitor_id"]]["status"] == "completion"
    }
    round_object["finish_order"] = {
        row["name"]: outcomes[row["local_competitor_id"]]["official_placing"]
        for row in rows
        if outcomes[row["local_competitor_id"]].get("official_placing") is not None
    }


def export_results(round_object, event, root_state, *, supersedes_settlement_id=None):
    """Append the signed settlement once, atomically, preserving the prior workbook."""
    from config import paths
    from woodchopping.data.excel_io import detect_results_sheet

    path = Path(paths.EXCEL_FILE)
    receipt_id = round_object["handicap_results"][0]["receipt_id"]
    settlement_id = round_object["v3_settlement_id"]
    backup_folder = (
        Path(os.environ["STRATHEX_WORKBOOK_BACKUP_DIR"]).resolve(strict=True)
        if os.environ.get("STRATHEX_WORKBOOK_BACKUP_DIR")
        else path.parent / "workbook-recovery"
    )
    if os.environ.get("STRATHEX_TEST_DB") != "1" and (
        not os.environ.get("STRATHEX_WORKBOOK_BACKUP_DIR") or backup_folder.stat().st_dev == path.stat().st_dev
    ):
        raise RuntimeError("official workbook export requires recovery on an independent filesystem")
    backup_folder.mkdir(mode=0o700, exist_ok=True)
    backup = backup_folder / f"{path.stem}-{datetime.now(timezone.utc).strftime('%Y%m%dT%H%M%S%fZ')}.xlsx"
    shutil.copy2(path, backup)
    if sha256(backup.read_bytes()).hexdigest() != sha256(path.read_bytes()).hexdigest():
        raise RuntimeError("workbook recovery copy did not verify; export stopped")
    workbook = load_workbook(path)
    try:
        sheet = detect_results_sheet(workbook)
        headers = [str(cell.value) if cell.value is not None else "" for cell in sheet[1]]
        required = [
            "CompetitorID",
            "Event",
            "Time (seconds)",
            "Size (mm)",
            "Species Code",
            "Quality",
            "HeatID",
            "Date",
            "Result Status",
            "Penalty (seconds)",
            "V3 Receipt",
            "V3 Settlement",
            "Result Revision",
            "Original Raw Time (seconds)",
            "Official Placing",
        ]
        date_header = "Date (optional)" if "Date (optional)" in headers and "Date" not in headers else "Date"
        required[required.index("Date")] = date_header
        for header in required:
            if header not in headers:
                headers.append(header)
                sheet.cell(1, len(headers), header)
        settlement_column = headers.index("V3 Settlement") + 1
        existing = [sheet.cell(index, settlement_column).value for index in range(2, sheet.max_row + 1)]
        if settlement_id in existing:
            return True
        if supersedes_settlement_id is not None:
            time_column = headers.index("Time (seconds)") + 1
            status_column = headers.index("Result Status") + 1
            raw_column = headers.index("Original Raw Time (seconds)") + 1
            for index in range(2, sheet.max_row + 1):
                if sheet.cell(index, settlement_column).value == supersedes_settlement_id:
                    original = sheet.cell(index, time_column).value
                    if isinstance(original, (int, float)):
                        sheet.cell(index, raw_column, original)
                    sheet.cell(index, time_column, "SUPERSEDED")
                    sheet.cell(index, status_column, "superseded_" + str(sheet.cell(index, status_column).value))
        for row in round_object["handicap_results"]:
            result = round_object["official_outcomes"][row["local_competitor_id"]]
            status = result["status"]
            values = {
                "CompetitorID": row["local_competitor_id"],
                "Event": str(event.get("event_code", root_state.get("event_code", "UH"))).upper(),
                "Time (seconds)": result["raw_time_ms"] / 1000 if status == "completion" else status.upper(),
                "Size (mm)": event.get("wood_diameter", root_state.get("wood_diameter")),
                "Species Code": event.get("wood_species", root_state.get("wood_species")),
                "Quality": event.get("wood_quality", root_state.get("wood_quality")),
                "HeatID": round_object.get("round_name", receipt_id),
                date_header: datetime.now(timezone.utc).replace(tzinfo=None),
                "Result Status": status,
                "Penalty (seconds)": result["penalty_ms"] / 1000 if result["penalty_ms"] is not None else None,
                "V3 Receipt": receipt_id,
                "V3 Settlement": settlement_id,
                "Result Revision": int(round_object.get("result_source_revision", 1)),
                "Official Placing": result.get("official_placing"),
                "Original Raw Time (seconds)": result["raw_time_ms"] / 1000
                if result["raw_time_ms"] is not None
                else None,
            }
            sheet.append([values.get(header) for header in headers])
        temporary = path.with_name(path.name + ".v3-next.xlsx")
        workbook.save(temporary)
        with temporary.open("rb+") as stream:
            os.fsync(stream.fileno())
        os.replace(temporary, path)
    finally:
        workbook.close()
    print("Signed V3 outcomes saved to the workbook. Nonfinishes and penalties remain outside raw-time history.")
    return True


def correct_saved_results(save_path, *, input_fn=input):
    """Deliberate operator correction against the saved exact Linux authority."""
    from woodchopping.strathmark_v3_client import build_v3_runtime
    from woodchopping.ui.handicap_ui import build_prediction_execution_context
    from woodchopping.ui.prediction_context import PredictionAuthorityStore
    from woodchopping.ui.state_persistence import load_tournament_state, save_tournament_state

    authority = PredictionAuthorityStore(os.environ["STRATHEX_PREDICTION_AUTHORITY_DB"])
    state = load_tournament_state(str(save_path), authority_store=authority)
    if state is None:
        raise RuntimeError("saved competition could not be verified")
    adapter = build_v3_runtime()
    if not hasattr(adapter, "correct_result"):
        raise RuntimeError("this competition needs the Linux lifecycle correction profile")
    context = build_prediction_execution_context(state, authority)
    rounds = []

    def visit(value):
        if isinstance(value, dict):
            if value.get("v3_settlement_id") and value.get("handicap_results"):
                rounds.append(value)
            for child in value.values():
                visit(child)
        elif isinstance(value, list):
            for child in value:
                visit(child)

    visit(state)
    for index, item in enumerate(rounds, 1):
        print(f"{index}. {item.get('round_name', 'Field')} | {item['v3_settlement_id']}")
    selected = int(input_fn("Correct which settled field? ")) - 1
    if not 0 <= selected < len(rounds):
        raise ValueError("select a listed settled field")
    item = rounds[selected]
    reason = input_fn("Official correction reason: ").strip()
    if not reason:
        raise ValueError("official correction reason is required")
    revised = {"handicap_results": item["handicap_results"]}
    if not collect_results(revised, input_fn=input_fn):
        return
    revision = int(item.get("result_source_revision", 1)) + 1
    receipt_id = item["handicap_results"][0]["receipt_id"]
    payload = {
        "schema_version": "strathmark-v3-settlement-request-v1",
        "receipt_id": receipt_id,
        "issue_batch_id": state["v3_issue_batches"][receipt_id],
        "observed_at_utc": datetime.now(timezone.utc).isoformat(timespec="milliseconds").replace("+00:00", "Z"),
        "deadline_ms": 10000,
        "supersedes_settlement_id": item["v3_settlement_id"],
        "reason_code": reason,
        "results": [
            {
                "competitor_id": row["competitor_id"],
                **revised["official_outcomes"][row["local_competitor_id"]],
                "source_revision": revision,
            }
            for row in item["handicap_results"]
        ],
    }
    pending = item.setdefault("v3_pending_correction", payload)
    if pending != payload:
        print("Resuming the previously saved correction with its original input.")
    if not save_tournament_state(state, str(save_path), authority_store=authority):
        raise RuntimeError("correction request could not be checkpointed")
    from woodchopping.ui.multi_event_ui import execute_with_v3_recovery

    response = execute_with_v3_recovery(
        lambda: adapter.correct_result(context, pending),
        root_state=state,
        authority_store=authority,
        v3_adapter=adapter,
        input_fn=input_fn,
    )
    if response is None:
        return
    state.setdefault("v3_result_corrections", []).append({"request": pending, "response": response})
    item["result_source_revision"] = response["source_revision"]
    item["v3_settlement_id"] = response["settlement_id"]
    item["official_outcomes"] = {
        row["local_competitor_id"]: {
            key: outcome[key] for key in ("status", "raw_time_ms", "penalty_ms", "official_placing")
        }
        for row, outcome in zip(item["handicap_results"], pending["results"], strict=True)
    }
    refresh_result_projections(item)
    if "prediction_comparison" in item:
        item.setdefault("superseded_prediction_comparisons", []).append(item.pop("prediction_comparison"))
    event = next(
        (
            event
            for event in state.get("events", [])
            if any(round_item is item for round_item in event.get("rounds", []))
        ),
        state,
    )
    export_results(item, event, state, supersedes_settlement_id=pending["supersedes_settlement_id"])
    item.pop("v3_pending_correction", None)
    if not save_tournament_state(state, str(save_path), authority_store=authority):
        raise RuntimeError("accepted correction could not be checkpointed")
    print(
        "Official correction retained as a new signed revision. Issued marks, frozen epochs, and saved advancement are preserved; review advancement separately."
    )
