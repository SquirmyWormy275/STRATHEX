"""Fail-closed protection for writes to the production Excel workbook.

Excel remains STRATHEX's judge-canonical record. A damaged, locked, or otherwise
unreadable existing workbook must never be replaced by a newly created partial
file. This module wraps the legacy results-entry routine without changing its
prompts or return contract.
"""

from __future__ import annotations

import copy
import hashlib
import os
import shutil
import tempfile
from types import ModuleType
from typing import Any, Callable

from openpyxl import load_workbook

_ROUND_STATE_KEYS = ("actual_results", "finish_order", "status")


def _required_sheet_names(excel_io_module: ModuleType) -> set[str]:
    paths = excel_io_module.paths
    return {
        str(paths.WOOD_SHEET).strip().lower(),
        str(paths.COMPETITOR_SHEET).strip().lower(),
        str(paths.RESULTS_SHEET).strip().lower(),
    }


def validate_complete_workbook(path: str, excel_io_module: ModuleType) -> None:
    """Raise when ``path`` is unreadable or lacks any required STRATHEX sheet."""
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        present = {str(name).strip().lower() for name in workbook.sheetnames}
    finally:
        workbook.close()

    missing = sorted(_required_sheet_names(excel_io_module) - present)
    if missing:
        raise ValueError("Workbook is incomplete; missing required sheet(s): " + ", ".join(missing))


def _file_digest(path: str) -> str:
    digest = hashlib.sha256()
    with open(path, "rb") as workbook_file:
        for chunk in iter(lambda: workbook_file.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _snapshot_round_state(round_object) -> dict[str, tuple[bool, Any]] | None:
    if not isinstance(round_object, dict):
        return None
    return {key: (key in round_object, copy.deepcopy(round_object.get(key))) for key in _ROUND_STATE_KEYS}


def _restore_round_state(
    round_object,
    snapshot: dict[str, tuple[bool, Any]] | None,
) -> None:
    if not isinstance(round_object, dict) or snapshot is None:
        return
    for key, (was_present, value) in snapshot.items():
        if was_present:
            round_object[key] = copy.deepcopy(value)
        else:
            round_object.pop(key, None)


def _actual_results_changed(
    round_object,
    snapshot: dict[str, tuple[bool, Any]] | None,
) -> bool:
    if not isinstance(round_object, dict) or snapshot is None:
        return False
    was_present, previous = snapshot["actual_results"]
    is_present = "actual_results" in round_object
    current = round_object.get("actual_results")
    return was_present != is_present or previous != current


def _create_backup(path: str) -> str:
    """Create a same-directory backup so restoration uses an atomic replace."""
    directory = os.path.dirname(os.path.abspath(path)) or os.curdir
    os.makedirs(directory, exist_ok=True)
    descriptor, backup_path = tempfile.mkstemp(
        prefix=f".{os.path.basename(path)}.",
        suffix=".prewrite-backup",
        dir=directory,
    )
    os.close(descriptor)
    try:
        shutil.copy2(path, backup_path)
    except Exception:
        try:
            os.remove(backup_path)
        except OSError:
            pass
        raise
    return backup_path


def _restore_backup(backup_path: str, target_path: str) -> bool:
    try:
        os.replace(backup_path, target_path)
        return True
    except OSError as exc:
        print("[X] CRITICAL: Workbook restoration failed.")
        print(f"    Backup retained at: {backup_path}")
        print(f"    Restore it manually to: {target_path}")
        print(f"    System error: {exc}")
        return False


def guarded_append_results_to_excel(
    append_function: Callable[..., Any],
    excel_io_module: ModuleType,
    *args,
    **kwargs,
):
    """Run the legacy append flow with fail-closed workbook protection.

    Behavior is intentionally conservative:

    1. A genuinely missing workbook is created through ``ensure_workbook()``,
       which restores the canonical seed or creates the complete three-sheet
       schema.
    2. An existing unreadable or incomplete workbook aborts the write before the
       interactive append routine can replace it.
    3. A same-directory backup is retained until the resulting workbook opens
       successfully and still contains Wood, Competitor, and Results sheets.
    4. The legacy routine's internal ``Workbook()`` fallback is disabled for the
       duration of the call, so a transient load failure cannot create a
       Results-only replacement.
    5. If tournament times were added in memory but the canonical workbook bytes
       did not change, only the affected round fields are rolled back.
    """
    target_path = str(excel_io_module.paths.EXCEL_FILE)

    try:
        if not os.path.exists(target_path):
            target_path = str(excel_io_module.ensure_workbook(target_path))
        validate_complete_workbook(target_path, excel_io_module)
    except Exception as exc:
        print("[X] Results were not written.")
        print(f"    Existing workbook could not be safely opened: {target_path}")
        print("    The file was left unchanged. Restore a valid workbook or close")
        print("    any program locking it, then retry.")
        print(f"    System error: {exc}")
        return None

    round_object = kwargs.get("round_object")
    round_snapshot = _snapshot_round_state(round_object)
    before_digest = _file_digest(target_path)

    try:
        backup_path = _create_backup(target_path)
    except Exception as exc:
        print("[X] Results were not written because a safety backup could not be created.")
        print(f"    Workbook left unchanged: {target_path}")
        print(f"    System error: {exc}")
        return None

    original_workbook_factory = excel_io_module.Workbook

    def _refuse_partial_replacement(*_args, **_kwargs):
        raise RuntimeError("Refusing to create a replacement workbook after an existing workbook load failure")

    excel_io_module.Workbook = _refuse_partial_replacement
    restore_attempted = False

    try:
        result = append_function(*args, **kwargs)

        try:
            validate_complete_workbook(target_path, excel_io_module)
        except Exception as exc:
            restore_attempted = True
            _restore_round_state(round_object, round_snapshot)
            restored = _restore_backup(backup_path, target_path)
            if restored:
                print("[X] Results write failed validation; the original workbook was restored.")
                print(f"    System error: {exc}")
            return None

        after_digest = _file_digest(target_path)
        if _actual_results_changed(round_object, round_snapshot) and after_digest == before_digest:
            _restore_round_state(round_object, round_snapshot)
            print("[X] Results were not written; tournament round state was restored.")
            return None

        return result
    except Exception as exc:
        restore_attempted = True
        _restore_round_state(round_object, round_snapshot)
        restored = _restore_backup(backup_path, target_path)
        if restored:
            print("[X] Results were not written; the original workbook was restored.")
            print(f"    System error: {exc}")
        return None
    finally:
        excel_io_module.Workbook = original_workbook_factory
        if not restore_attempted and os.path.exists(backup_path):
            try:
                os.remove(backup_path)
            except OSError:
                # A retained backup is harmless and preferable to deleting a file
                # whose cleanup status is uncertain.
                pass
