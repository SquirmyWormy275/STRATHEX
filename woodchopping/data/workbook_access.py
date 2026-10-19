"""Bounded workbook snapshots and private, atomic canonical writes."""

import os
import stat
from pathlib import Path
from uuid import uuid4

from openpyxl import load_workbook as _load_workbook
from strathmark.input_limits import bounded_workbook
from strathmark.private_storage import verify_authority_parent


def _identity(status):
    return status.st_dev, status.st_ino, status.st_size, status.st_mtime_ns


def load_workbook(path, **kwargs):
    """Parse exactly the bounded snapshot, retaining its identity for publication."""
    path = Path(path).absolute()
    verify_authority_parent(path)
    before = path.lstat()
    if not stat.S_ISREG(before.st_mode) or before.st_nlink != 1:
        raise ValueError("Workbook must be a regular file with one link")
    snapshot = bounded_workbook(path)
    if _identity(path.lstat()) != _identity(before):
        raise RuntimeError("workbook changed while reading")
    workbook = _load_workbook(snapshot, **kwargs)
    workbook._source_status = before
    return workbook


def save_workbook(workbook, path, original_status=None):
    """Publish only the exclusively created file we saved and synced."""
    from strathmark.private_storage import open_private_exclusive, verify_authority_parent

    path = Path(path).absolute()
    verify_authority_parent(path)
    temporary = path.with_name(f".{path.name}.{uuid4().hex}.tmp")
    descriptor = open_private_exclusive(temporary)
    try:
        with os.fdopen(descriptor, "w+b") as stream:
            workbook.save(stream)
            stream.flush()
            os.fsync(stream.fileno())
            saved = os.fstat(stream.fileno())
            staged = temporary.lstat()
            if (
                not stat.S_ISREG(staged.st_mode)
                or staged.st_nlink != 1
                or (staged.st_dev, staged.st_ino) != (saved.st_dev, saved.st_ino)
            ):
                raise RuntimeError("workbook staging file changed; export stopped")
            verify_authority_parent(path)
            expected = original_status if original_status is not None else getattr(workbook, "_source_status", None)
            try:
                current = path.lstat()
            except FileNotFoundError:
                current = None
            if (current is None) != (expected is None) or (
                current is not None
                and (
                    not stat.S_ISREG(current.st_mode)
                    or current.st_nlink != 1
                    or _identity(current) != _identity(expected)
                )
            ):
                raise RuntimeError("workbook changed during export; retry against the current workbook")
        os.replace(temporary, path)
        if os.name == "posix":
            directory = os.open(path.parent, os.O_RDONLY | os.O_DIRECTORY)
            try:
                os.fsync(directory)
            finally:
                os.close(directory)
    finally:
        temporary.unlink(missing_ok=True)
