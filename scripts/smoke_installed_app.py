"""Install an exact wheel and drive the terminal app against synthetic data."""

from __future__ import annotations

import argparse
import os
import subprocess
import tarfile
import tempfile
import venv
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--wheel", type=Path, required=True)
    parser.add_argument("--sdist", type=Path)
    args = parser.parse_args()
    wheel = args.wheel.resolve(strict=True)
    if args.sdist:
        with tarfile.open(args.sdist.resolve(strict=True)) as archive:
            if any(
                Path(item.name).name in {"woodchopping.xlsx", "woodchopping_clean.xlsx", "installation-key.pem"}
                for item in archive.getmembers()
            ):
                raise ValueError("operator workbook or signing key must not enter a public source release")
    with tempfile.TemporaryDirectory(prefix="strathex-installed-") as temporary:
        root = Path(temporary)
        environment = root / "environment"
        venv.EnvBuilder(with_pip=True).create(environment)
        python = environment / ("Scripts/python.exe" if os.name == "nt" else "bin/python")
        subprocess.run(
            [str(python), "-m", "pip", "install", "--disable-pip-version-check", str(wheel)], check=True, timeout=300
        )
        workbook = Workbook()
        workbook.remove(workbook.active)
        roster = workbook.create_sheet("Competitor")
        roster.append(["CompetitorID", "Name", "Country", "State/Province", "Gender"])
        roster.append(["SYN001", "Synthetic One", "SYN", "", "M"])
        workbook.create_sheet("wood").append(
            [
                "Scientific Name",
                "species",
                "speciesID",
                "country",
                "region",
                "janka_hard",
                "spec_gravity",
                "crush_strength",
                "shear",
                "MOR",
                "MOE",
            ]
        )
        history = workbook.create_sheet("Results")
        history.append(["CompetitorID", "Event", "Time (seconds)", "Species Code", "Size (mm)", "Quality", "Date"])
        history.append(["SYN001", "UH", 30, "S01", 300, 5, datetime(2024, 1, 1)])
        history.append(["SYN001", "UH", 31, "S01", 300, 5, datetime(2025, 1, 1)])
        workbook_path = root / "synthetic.xlsx"
        workbook.save(workbook_path)
        workbook.close()
        child_env = {
            key: value for key, value in os.environ.items() if not key.startswith(("STRATHMARK_", "STRATHEX_"))
        }
        child_env.update(STRATHMARK_TEST_DB="1", PYTHONUTF8="1")
        completed = subprocess.run(
            [str(python), "-m", "strathex_cli", "--workbook", str(workbook_path), "--data-dir", str(root / "data")],
            input="9\nn\n",
            cwd=root,
            env=child_env,
            capture_output=True,
            text=True,
            encoding="utf-8",
            timeout=60,
        )
        if (
            completed.returncode
            or "Goodbye!" not in completed.stdout
            or "Error loading results" in completed.stdout
            or "migrated 2 historical results" not in completed.stdout
        ):
            raise RuntimeError(completed.stdout + completed.stderr)
        print("Installed STRATHEX started and exited successfully with synthetic data outside the checkout.")


if __name__ == "__main__":
    main()
