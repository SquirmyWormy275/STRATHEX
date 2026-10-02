"""Run installed V2/V3 choice, real V3 numerics, saved restart, and no-issue checks.

Run with the installed STRATHEX interpreter and a separate installed V3 Python.
All workbook rows, rosters, events, and databases here are synthetic. No runtime
test fixtures, presealed forecasts, or mocked assessor outputs are used.
"""

from __future__ import annotations

import argparse
import json
import os
import subprocess
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--v3-python", type=Path, required=True)
    parser.add_argument("--ml-bundle", type=Path)
    parser.add_argument("--training-source", help="If no bundle is supplied, train on this synthetic workbook only")
    parser.add_argument("--training-repository", type=Path)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    root = args.output.resolve()
    root.mkdir(parents=True, exist_ok=False)
    workbook = Workbook()
    workbook.remove(workbook.active)
    competitor = workbook.create_sheet("Competitor")
    competitor.append(["CompetitorID", "Name", "Country", "State/Province", "Gender"])
    competitor.append(["SYN001", "Synthetic One", "SYN", "Synthetic", "M"])
    competitor.append(["SYN002", "Synthetic Two", "SYN", "Synthetic", "M"])
    wood = workbook.create_sheet("Wood")
    wood.append(["speciesID", "species", "spec_gravity"])
    wood.append(["S01", "Synthetic wood", 0.5])
    history = workbook.create_sheet("Results")
    history.append(
        [
            "CompetitorID",
            "Event",
            "Time (seconds)",
            "Size (mm)",
            "Species Code",
            "Date (optional)",
            "Notes (Competition, special circumstances, etc.)",
        ]
    )
    for year in range(2021, 2026):
        for competitor_id, seconds in (("SYN001", 28), ("SYN002", 40)):
            history.append([competitor_id, "UH", seconds, 300, "S01", datetime(year, 1, 1), f"Synthetic {year}"])
    workbook_path = root / "synthetic.xlsx"
    workbook.save(workbook_path)
    workbook.close()
    os.environ.update(
        STRATHEX_WORKBOOK=str(workbook_path),
        STRATHMARK_TEST_DB="1",
        STRATHMARK_DB_PATH=str(root / "v2-results.db"),
        STRATHEX_PREDICTION_AUTHORITY_DB=str(root / "authority.db"),
        STRATHEX_V3_COMMAND_DB=str(root / "commands.db"),
    )
    # Select disposable paths before importing any application module.
    from woodchopping.strathmark_v3_client import V3ClientError
    from woodchopping.strathmark_v3_local import LocalV3Candidate
    from woodchopping.ui.handicap_ui import (
        build_engine_router,
        calculate_authoritative_field,
        calculate_authoritative_seeding,
    )
    from woodchopping.ui.multi_event_ui import select_prediction_engine_for_scope
    from woodchopping.ui.prediction_context import PredictionAuthorityStore, resolve_authority_for_state
    from woodchopping.ui.state_persistence import load_tournament_state, save_tournament_state

    if args.ml_bundle is None:
        if not args.training_source or not args.training_repository:
            parser.error("supply --ml-bundle or --training-source for synthetic-only training")
        subprocess.run(
            [
                str(args.v3_python.absolute()),
                "-I",
                "-m",
                "strathmark.v3.factory.candidate_cli",
                "--workbook",
                str(workbook_path),
                "--output",
                str(root / "synthetic-model"),
                "--source-commit",
                args.training_source,
                "--source-repository",
                str(args.training_repository.resolve(strict=True)),
                "--cutoff-at-utc",
                "2026-10-02T00:00:00.000Z",
            ],
            check=True,
            timeout=120,
            stdout=subprocess.DEVNULL,
        )
        args.ml_bundle = root / "synthetic-model/ml-bundle"
    candidate = LocalV3Candidate(
        python=args.v3_python, ml_bundle=args.ml_bundle, workbook=workbook_path, snapshot_root=root / "snapshots"
    )
    store = PredictionAuthorityStore(root / "authority.db")
    roster = pd.DataFrame(
        {"competitor_id": ["SYN001", "SYN002"], "competitor_name": ["Synthetic One", "Synthetic Two"]}
    )
    history_frame = pd.DataFrame(
        [
            {
                "competitor_id": competitor_id,
                "competitor_name": name,
                "event": "UH",
                "raw_time": seconds,
                "species": "S01",
                "size_mm": 300,
                "quality": 5,
                "date": f"{year}-01-01",
            }
            for year in range(2022, 2026)
            for competitor_id, name, seconds in (("SYN001", "Synthetic One", 28), ("SYN002", "Synthetic Two", 40))
        ]
    )
    router = build_engine_router(v3_adapter=candidate)

    def select(engine):
        state = {}
        answers = iter(["1" if engine == "v2" else "2", "synthetic_verification", ""])
        select_prediction_engine_for_scope(
            state,
            authority_store=store,
            owner_kind="single_event",
            readiness_provider=candidate.selector_readiness,
            input_fn=lambda _: next(answers),
        )
        return state

    states = {engine: select(engine) for engine in ("v2", "v3")}

    def request(state, authority=store, selected_router=router):
        save = root / f"{resolve_authority_for_state(state, authority).engine}-event.json"
        return dict(
            root_state=state,
            authority_store=authority,
            engine_router=selected_router,
            field_local_id="round-one:heat-one",
            competitors_df=roster,
            wood_species="S01",
            wood_diameter=300,
            wood_quality=5,
            event_code="UH",
            results_df=history_frame,
            prediction_as_of="2026-10-02",
            include_store_history=False,
            checkpoint_callback=lambda value: save_tournament_state(value, str(save), authority_store=authority),
        )

    v2 = calculate_authoritative_field(**request(states["v2"]))
    assert all(row["engine_version"].startswith("2.") for row in v2)
    seeding = calculate_authoritative_seeding(forecast_adapter=candidate.forecast_seeding, **request(states["v3"]))
    assert all("mark" not in row and row["predicted_time"] > 0 for row in seeding)
    v3 = calculate_authoritative_field(**request(states["v3"]))
    assert all(row["engine_version"].startswith("3.") and row["std_dev"] > 0 for row in v3)
    assert v3[0]["predicted_time"] < v3[1]["predicted_time"]
    assert v3[0]["mark"] > v3[1]["mark"] == 3
    states["v3"].update(
        handicap_results_all=v3,
        all_competitors_df=roster,
        all_competitors=roster["competitor_name"].tolist(),
        prediction_as_of="2026-10-02",
    )
    assert save_tournament_state(states["v3"], str(root / "v3-event.json"), authority_store=store)
    reopened = PredictionAuthorityStore(root / "authority.db")
    resumed = load_tournament_state(str(root / "v3-event.json"), authority_store=reopened)
    assert resumed is not None and resolve_authority_for_state(resumed, reopened).engine == "v3"
    assert resumed["handicap_results_all"] == v3
    restarted = LocalV3Candidate(
        python=args.v3_python, ml_bundle=args.ml_bundle, workbook=workbook_path, snapshot_root=root / "snapshots"
    )
    repeated = calculate_authoritative_field(**request(resumed, reopened, build_engine_router(v3_adapter=restarted)))
    assert repeated == v3
    try:
        restarted.acknowledge_issue()
    except V3ClientError:
        pass
    else:
        raise AssertionError("numeric candidate authorized issue")
    (root / "report.json").write_text(
        json.dumps(
            {
                "synthetic_only": True,
                "mocked_numerics": False,
                "v2": v2,
                "v3": v3,
                "pre_field_has_marks": False,
                "restart_same_output": True,
                "issue_blocked": True,
                "readiness": candidate.selector_readiness(),
            },
            indent=2,
        ),
        encoding="utf-8",
    )
    print("Installed V2/V3 selection, real Formula + ML numerics, restart, and no-issue checks passed.")


if __name__ == "__main__":
    main()
