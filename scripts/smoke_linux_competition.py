"""Installed full Linux workflow with synthetic inputs and actual native ML inference."""

from __future__ import annotations

import argparse
import json
import os
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--v3-python", type=Path, required=True)
    parser.add_argument("--ml-bundle", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--backup-dir", type=Path)
    args = parser.parse_args()
    root = args.output.resolve()
    root.mkdir(parents=True, exist_ok=False)
    workbook = Workbook()
    workbook.remove(workbook.active)
    competitor = workbook.create_sheet("Competitor")
    competitor.append(["CompetitorID", "Name", "Country", "State/Province", "Gender"])
    competitor.append(["SYN001", "Synthetic One", "SYN", "Synthetic", "M"])
    competitor.append(["SYN002", "Synthetic Two", "SYN", "Synthetic", "M"])
    wood = workbook.create_sheet("Wood")
    wood.append(["speciesID", "species", "spec_gravity"])
    wood.append(["S01", "Synthetic wood", 0.5])
    history = workbook.create_sheet("Results")
    history.append(["CompetitorID", "Event", "Time (seconds)", "Size (mm)", "Species Code", "Date (optional)"])
    for year in range(2021, 2026):
        for identifier, seconds in (("SYN001", 28), ("SYN002", 40)):
            history.append([identifier, "UH", seconds, 300, "S01", datetime(year, 1, 1)])
    workbook_path = root / "synthetic.xlsx"
    workbook.save(workbook_path)
    workbook.close()
    os.environ.update(
        STRATHEX_WORKBOOK=str(workbook_path),
        STRATHMARK_TEST_DB="1",
        STRATHMARK_DB_PATH=str(root / "v2-results.db"),
        STRATHEX_PREDICTION_AUTHORITY_DB=str(root / "authority.db"),
        STRATHEX_V3_COMMAND_DB=str(root / "commands.db"),
    )
    from woodchopping.strathmark_v3_linux import LinuxV3Competition
    from woodchopping.ui.handicap_ui import (
        build_engine_router,
        calculate_authoritative_field,
        calculate_authoritative_seeding,
    )
    from woodchopping.ui.multi_event_ui import (
        finalize_completed_competition,
        record_and_settle_v3_round,
        review_v3_approval_queue,
        select_prediction_engine_for_scope,
    )
    from woodchopping.ui.prediction_context import PredictionAuthorityStore, resolve_authority_for_state
    from woodchopping.ui.schedule_printout import generate_printable_schedule
    from woodchopping.ui.state_persistence import load_tournament_state, save_tournament_state
    from woodchopping.v3_authority_store import V3CommandStore

    authority = PredictionAuthorityStore(root / "authority.db")

    def runtime():
        return LinuxV3Competition(
            python=args.v3_python,
            ml_bundle=args.ml_bundle,
            workbook=workbook_path,
            runtime_root=root / "installation",
            command_store=V3CommandStore(root / "commands.db"),
            backup_dir=args.backup_dir,
        )

    adapter = runtime()
    router = build_engine_router(v3_adapter=adapter)

    def selected(engine):
        state = {}
        answers = iter(["2" if engine == "v3" else "1", "synthetic_verification", ""])
        select_prediction_engine_for_scope(
            state,
            authority_store=authority,
            owner_kind="single_event",
            readiness_provider=adapter.selector_readiness,
            input_fn=lambda _: next(answers),
        )
        return state

    state, v2_state = selected("v3"), selected("v2")
    roster = pd.DataFrame(
        {"competitor_id": ["SYN001", "SYN002"], "competitor_name": ["Synthetic One", "Synthetic Two"]}
    )
    history_frame = pd.DataFrame(
        [
            {
                "competitor_id": identifier,
                "competitor_name": name,
                "event": "UH",
                "raw_time": seconds,
                "species": "S01",
                "size_mm": 300,
                "quality": 5,
                "date": f"{year}-01-01",
            }
            for year in range(2022, 2026)
            for identifier, name, seconds in (("SYN001", "Synthetic One", 28), ("SYN002", "Synthetic Two", 40))
        ]
    )
    state.update(
        all_competitors_df=roster,
        all_competitors=roster["competitor_name"].tolist(),
        event_name="Synthetic event",
        wood_species="S01",
        wood_diameter=300,
        wood_quality=5,
        event_code="UH",
        rounds=[],
        prediction_as_of="2026-10-02",
    )

    def checkpoint(value):
        return save_tournament_state(value, str(root / "competition.json"), authority_store=authority)

    def request(owner, name, ordinal=1):
        return dict(
            root_state=owner,
            authority_store=authority,
            engine_router=router,
            field_local_id=name,
            round_local_id=f"stage-{ordinal}",
            round_ordinal=ordinal,
            competitors_df=roster,
            wood_species="S01",
            wood_diameter=300,
            wood_quality=5,
            event_code="UH",
            results_df=history_frame,
            prediction_as_of="2026-10-02",
            include_store_history=False,
            checkpoint_callback=checkpoint,
        )

    baseline = calculate_authoritative_field(**request(v2_state, "v2-field"))
    assert all(row["engine_version"].startswith("2.") for row in baseline)
    prefield = calculate_authoritative_seeding(forecast_adapter=adapter.forecast_seeding, **request(state, "seeding"))
    assert all("mark" not in row for row in prefield)
    first = calculate_authoritative_field(**request(state, "heat-1"))
    assert first[0]["predicted_time"] < first[1]["predicted_time"] and first[0]["mark"] > first[1]["mark"] == 3

    def add_round(rows, name):
        item = {
            "round_name": name,
            "round_type": "heat" if name.startswith("Heat") else "final",
            "status": "pending",
            "competitors": roster["competitor_name"].tolist(),
            "handicap_results": rows,
            "v3_round_id": rows[0]["round_id"],
            "num_to_advance": 2,
            "actual_results": {},
            "finish_order": {},
        }
        state["rounds"].append(item)
        return item

    heat = add_round(first, "Heat 1")
    # Unissued sheets and result entry must fail without changing the workbook.
    before = workbook_path.read_bytes()
    blocked = record_and_settle_v3_round(
        state,
        state,
        heat,
        write_action=lambda: (_ for _ in ()).throw(AssertionError("unissued write")),
        authority_store=authority,
        v3_adapter=adapter,
        input_fn=lambda _: "c",
    )
    assert blocked is False and workbook_path.read_bytes() == before
    try:
        generate_printable_schedule(state, authority_store=authority)
    except (ValueError, RuntimeError):
        pass
    else:
        raise AssertionError("unissued field printed")

    def review(answer_issue="y"):
        decisions = review_v3_approval_queue(
            state,
            authority_store=authority,
            v3_adapter=adapter,
            input_fn=lambda prompt: answer_issue if "Issue these" in prompt else "y" if "batch" in prompt else "a",
            checkpoint_callback=checkpoint,
        )
        assert decisions

    review("n")
    assert state["v3_issue_status"] == "approved_unissued" and not state.get("v3_issue_batches")
    review("y")
    assert heat["handicap_results"][0]["issue_batch_id"]

    def settle(item, outcomes):
        answers = iter([*outcomes, "y"])
        assert record_and_settle_v3_round(
            state,
            state,
            item,
            write_action=lambda: (_ for _ in ()).throw(AssertionError("legacy V3 write")),
            authority_store=authority,
            v3_adapter=adapter,
            input_fn=lambda _: next(answers),
        )
        item["status"] = "completed"
        assert checkpoint(state)

    settle(heat, ["22", "50"])
    same_round = calculate_authoritative_field(**request(state, "heat-2"))
    assert [(row["predicted_time"], row["mark"], row["epoch_digest"]) for row in same_round] == [
        (row["predicted_time"], row["mark"], row["epoch_digest"]) for row in first
    ]
    second_heat = add_round(same_round, "Heat 2")
    review()
    settle(second_heat, ["23", "DNS"])
    saved = load_tournament_state(str(root / "competition.json"), authority_store=authority)
    assert saved is not None and resolve_authority_for_state(saved, authority).engine == "v3"
    state = saved
    adapter = runtime()
    router = build_engine_router(v3_adapter=adapter)
    recovered = calculate_authoritative_field(**request(state, "heat-1"))
    assert [(row["predicted_time"], row["mark"], row["receipt_id"]) for row in recovered] == [
        (row["predicted_time"], row["mark"], row["receipt_id"]) for row in first
    ]
    final = calculate_authoritative_field(field_kind="championship", **request(state, "final", 2))
    assert {row["mark"] for row in final} == {3}
    assert (
        final[0]["epoch_digest"] != first[0]["epoch_digest"]
        and final[0]["predicted_time"] != first[0]["predicted_time"]
    )
    final_round = add_round(final, "Final")
    review()
    settle(final_round, ["24", "PENALTY 46 2"])
    state["final_results"] = {"winner": "Synthetic One"}
    finalize_completed_competition(state, authority_store=authority, v3_adapter=adapter, prompt_for_feedback=False)
    assert state["prediction_engine_closure"]["status"] == "closed"
    assert checkpoint(state)
    os.environ.update(
        STRATHEX_V3_LOCAL_PYTHON=str(args.v3_python.absolute()),
        STRATHEX_V3_LOCAL_ML_BUNDLE=str(args.ml_bundle.absolute()),
        STRATHEX_V3_LOCAL_RUNTIME_ROOT=str(root / "installation"),
    )
    from woodchopping.ui.linux_results import correct_saved_results

    answers = iter(["3", "official synthetic correction", "25", "DNS", "y"])
    correct_saved_results(root / "competition.json", input_fn=lambda _: next(answers))
    corrected = load_tournament_state(str(root / "competition.json"), authority_store=authority)
    assert corrected["v3_result_corrections"][0]["response"]["source_revision"] == 2
    assert corrected["rounds"][-1]["handicap_results"] == state["rounds"][-1]["handicap_results"]
    assert corrected["final_results"] == state["final_results"]
    workbook = __import__("openpyxl").load_workbook(workbook_path, read_only=True)
    headers = [cell.value for cell in workbook["Results"][1]]
    assert "superseded_completion" in [
        row[headers.index("Result Status")] for row in workbook["Results"].iter_rows(min_row=2, values_only=True)
    ]
    workbook.close()
    (root / "report.json").write_text(
        json.dumps(
            {
                "synthetic_only": True,
                "mocked_numerics": False,
                "v2_available": True,
                "prefield_mark_free": True,
                "approval_separate_from_issue": True,
                "unissued_print_and_results_blocked": True,
                "restart_replayed_same_receipt": True,
                "same_round_frozen": True,
                "next_round_changed": True,
                "explicit_nonfinishes_and_penalties": True,
                "competition_closed": True,
                "championship_mark_3_receipts": True,
                "signed_corrections_preserve_epochs_and_advancement": True,
                "closure": state["prediction_engine_closure"],
                "v2": baseline,
                "first": first,
                "final": final,
            },
            indent=2,
        )
    )
    print(
        "Installed Linux competition completed: V2/V3 choice, real Formula+ML, approval, issue, results, restart, later-round learning, closure."
    )


if __name__ == "__main__":
    main()
