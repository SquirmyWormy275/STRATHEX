"""Round-state rollback tests for failed canonical workbook writes."""

from __future__ import annotations

from pathlib import Path
from types import SimpleNamespace

from openpyxl import Workbook, load_workbook

import woodchopping.data.excel_io as excel_io
from woodchopping.data.workbook_guard import guarded_append_results_to_excel


def _paths(path: Path) -> SimpleNamespace:
    return SimpleNamespace(
        EXCEL_FILE=str(path),
        WOOD_SHEET="wood",
        COMPETITOR_SHEET="Competitor",
        RESULTS_SHEET="Results",
    )


def _write_complete_workbook(path: Path) -> None:
    workbook = Workbook()
    workbook.remove(workbook["Sheet"])
    workbook.create_sheet("wood").append(excel_io.WOOD_HEADERS)
    workbook.create_sheet("Competitor").append(excel_io.COMPETITOR_HEADERS)
    workbook.create_sheet("Results").append(
        [
            "CompetitorID",
            "Event",
            "Time (seconds)",
            "Size (mm)",
            "Species Code",
            "Quality",
            "HeatID",
            "Date",
        ]
    )
    workbook.save(path)
    workbook.close()


def test_round_fields_roll_back_when_times_change_but_workbook_does_not(
    tmp_path,
    monkeypatch,
    capsys,
):
    path = tmp_path / "woodchopping.xlsx"
    _write_complete_workbook(path)
    monkeypatch.setattr(excel_io, "paths", _paths(path))

    round_object = {
        "actual_results": {"Existing Axe": 31.0},
        "finish_order": {"Existing Axe": 1},
        "status": "pending",
        "round_name": "Heat 1",
    }

    def failed_append(*_args, **_kwargs):
        round_object["actual_results"]["Alice Axe"] = 30.0
        round_object["finish_order"]["Alice Axe"] = 2
        round_object["status"] = "in_progress"
        # Simulate the legacy function swallowing its own workbook-save error.
        return None

    result = guarded_append_results_to_excel(
        failed_append,
        excel_io,
        round_object=round_object,
    )

    assert result is None
    assert round_object == {
        "actual_results": {"Existing Axe": 31.0},
        "finish_order": {"Existing Axe": 1},
        "status": "pending",
        "round_name": "Heat 1",
    }
    assert "tournament round state was restored" in capsys.readouterr().out


def test_round_fields_are_retained_when_workbook_write_succeeds(tmp_path, monkeypatch):
    path = tmp_path / "woodchopping.xlsx"
    _write_complete_workbook(path)
    monkeypatch.setattr(excel_io, "paths", _paths(path))

    round_object = {
        "actual_results": {},
        "finish_order": {},
        "status": "pending",
    }

    def successful_append(*_args, **_kwargs):
        round_object["actual_results"]["Alice Axe"] = 30.0
        round_object["finish_order"]["Alice Axe"] = 1
        round_object["status"] = "in_progress"

        workbook = load_workbook(path)
        workbook["Results"].append(
            ["C001", "SB", 30.0, 300, "S01", 5, "H1", "2026-08-18"]
        )
        workbook.save(path)
        workbook.close()
        return "saved"

    result = guarded_append_results_to_excel(
        successful_append,
        excel_io,
        round_object=round_object,
    )

    assert result == "saved"
    assert round_object == {
        "actual_results": {"Alice Axe": 30.0},
        "finish_order": {"Alice Axe": 1},
        "status": "in_progress",
    }
