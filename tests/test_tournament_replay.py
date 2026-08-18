"""Deterministic replay coverage for judge-facing tournament workflows.

Every persistence boundary in this module is redirected to ``tmp_path``.  The
tests never open the repository's production workbook or default ResultStore.
"""

from __future__ import annotations

from pathlib import Path
from types import SimpleNamespace

import pandas as pd
from openpyxl import Workbook, load_workbook

import woodchopping.data.excel_io as excel_io
import woodchopping.data.store_registry as store_registry
from woodchopping.data import append_results_to_excel
from woodchopping.ui.bracket_ui import (
    generate_bracket_with_byes,
    get_current_match,
    initialize_bracket_tournament,
    record_match_result,
)
from woodchopping.ui.multi_event_ui import (
    complete_event_round,
    generate_complete_day_schedule,
    generate_tournament_summary,
    get_next_incomplete_round,
    sequential_results_workflow,
)
from woodchopping.ui.payout_ui import display_single_event_final_results
from woodchopping.ui.state_persistence import (
    load_multi_event_tournament,
    load_tournament_state,
    save_multi_event_tournament,
    save_tournament_state,
)
from woodchopping.ui.tournament_ui import (
    complete_recorded_round,
    current_stage_rounds,
    distribute_competitors_into_heats,
    generate_next_round,
    select_heat_advancers,
)


class RecordingStore:
    """Minimal isolated ResultStore substitute used to verify dual writes."""

    def __init__(self) -> None:
        self.rows: list[dict] = []

    def record_result(self, **row) -> bool:
        self.rows.append(row)
        return True


def _competitors(names: list[str]) -> pd.DataFrame:
    return pd.DataFrame(
        {
            "competitor_name": names,
            "CompetitorID": [f"C{i:03d}" for i in range(1, len(names) + 1)],
            "competitor_country": ["USA"] * len(names),
            "gender": ["X"] * len(names),
        }
    )


def _handicaps(names: list[str], mark: int = 3) -> list[dict]:
    return [
        {
            "name": name,
            "mark": mark,
            "predicted_time": 30.0 + index,
            "method_used": "Replay fixture",
            "confidence": "test",
            "predictions": {},
        }
        for index, name in enumerate(names)
    ]


def _write_workbook(path: Path, names: list[str]) -> None:
    workbook = Workbook()
    workbook.remove(workbook["Sheet"])
    workbook.create_sheet("wood").append(excel_io.WOOD_HEADERS)
    competitors = workbook.create_sheet("Competitor")
    competitors.append(excel_io.COMPETITOR_HEADERS)
    for index, name in enumerate(names, 1):
        competitors.append([f"C{index:03d}", name, "USA", "MT", "X"])
    excel_io.detect_results_sheet(workbook)
    workbook.save(path)
    workbook.close()


def _script_input(monkeypatch, answers: list[str]) -> None:
    scripted = iter(answers)
    monkeypatch.setattr("builtins.input", lambda _prompt="": next(scripted))


def test_single_event_replay_records_resumes_and_completes(tmp_path, monkeypatch):
    names = ["Alice Axe", "Bob Block", "Carol Chop", "Drew Drive"]
    workbook_path = tmp_path / "woodchopping.xlsx"
    state_path = tmp_path / "tournament_state.json"
    _write_workbook(workbook_path, names)
    monkeypatch.setattr(
        excel_io,
        "paths",
        SimpleNamespace(
            EXCEL_FILE=str(workbook_path),
            WOOD_SHEET="wood",
            COMPETITOR_SHEET="Competitor",
            RESULTS_SHEET="Results",
        ),
    )
    store = RecordingStore()
    monkeypatch.setattr(store_registry, "get_store", lambda: store)

    roster = _competitors(names)
    heats = distribute_competitors_into_heats(roster, _handicaps(names, mark=7), 2, 2)
    for heat in heats:
        heat["num_to_advance"] = 1
    state = {
        "event_name": "Replay Cup",
        "format": "heats_to_finals",
        "num_stands": 2,
        "all_competitors": names,
        "all_competitors_df": roster,
        "rounds": heats,
    }
    wood = {"species": "S01", "size_mm": 300, "quality": 5, "event": "SB"}

    first_heat = state["rounds"][0]
    _script_input(monkeypatch, ["2", "30", "33", ""])
    assert append_results_to_excel(None, wood, round_object=first_heat, tournament_state=state)
    select_heat_advancers(first_heat)
    save_tournament_state(state, str(state_path))
    state = load_tournament_state(str(state_path))
    assert state is not None

    second_heat = state["rounds"][1]
    _script_input(monkeypatch, ["2", "31", "34", ""])
    assert append_results_to_excel(None, wood, round_object=second_heat, tournament_state=state)
    select_heat_advancers(second_heat)

    stage_type, stage_rounds = current_stage_rounds(state["rounds"])
    assert stage_type == "heat"
    advancers = [name for heat in stage_rounds for name in heat["advancers"]]
    final = generate_next_round(state, advancers, "final", is_championship=True)[0]
    state["rounds"].append(final)

    _script_input(monkeypatch, ["2", "29", "31"])
    write_succeeded = append_results_to_excel(None, wood, round_object=final, tournament_state=state)
    assert complete_recorded_round(state, final, write_succeeded)
    save_tournament_state(state, str(state_path))
    reloaded = load_tournament_state(str(state_path))

    assert reloaded is not None
    assert reloaded["rounds"][-1]["status"] == "completed"
    assert all(item["mark"] == 3 for item in reloaded["rounds"][-1]["handicap_results"])
    workbook = load_workbook(workbook_path, read_only=True)
    assert workbook["Results"].max_row == 7  # header + four heats + two finalists
    workbook.close()
    assert len(store.rows) == 6
    competition_ids = {row["competition_id"] for row in store.rows}
    assert competition_ids == {reloaded["competition_id"]}
    assert reloaded["competition_id"].startswith("strathex:")


def test_failed_result_entry_keeps_terminal_round_retryable():
    final = {
        "round_name": "Final",
        "round_type": "final",
        "competitors": ["Alice", "Bob"],
        "competitors_df": _competitors(["Alice", "Bob"]),
        "status": "pending",
        "finish_order": {},
    }
    state = {"format": "heats_to_finals", "rounds": [final]}

    assert not complete_recorded_round(state, final, entry_succeeded=False)
    assert final["status"] == "pending"
    assert "final_results" not in state


def test_current_stage_ignores_completed_prior_rounds():
    rounds = [
        {"round_type": "heat", "status": "completed", "advancers": ["Alice"]},
        {"round_type": "heat", "status": "completed", "advancers": ["Bob"]},
        {"round_type": "semi", "status": "completed", "advancers": ["Carol"]},
        {"round_type": "semi", "status": "completed", "advancers": ["Drew"]},
    ]

    stage_type, stage_rounds = current_stage_rounds(rounds)

    assert stage_type == "semi"
    assert [round_object["advancers"] for round_object in stage_rounds] == [["Carol"], ["Drew"]]


def test_multi_event_single_heat_completes_and_resumes(tmp_path):
    round_object = {
        "round_name": "Heat 1",
        "round_type": "heat",
        "competitors": ["Alice", "Bob"],
        "competitors_df": _competitors(["Alice", "Bob"]),
        "status": "in_progress",
        "finish_order": {"Alice": 1, "Bob": 2},
    }
    event = {
        "event_name": "300mm SB",
        "event_type": "championship",
        "format": "single_heat",
        "status": "in_progress",
        "all_competitors": ["Alice", "Bob"],
        "all_competitors_df": _competitors(["Alice", "Bob"]),
        "rounds": [round_object],
    }
    tournament = {
        "tournament_name": "Replay Day",
        "events": [event],
        "total_events": 1,
        "events_completed": 0,
        "current_event_index": 0,
    }

    complete_event_round(tournament, event, round_object)
    path = tmp_path / "multi_tournament_state.json"
    save_multi_event_tournament(tournament, str(path))
    reloaded = load_multi_event_tournament(str(path))

    assert reloaded is not None
    assert reloaded["events_completed"] == 1
    assert reloaded["events"][0]["status"] == "completed"
    assert reloaded["events"][0]["final_results"]["first_place"] == "Alice"


def test_event_jump_changes_the_next_round_selected():
    tournament = {
        "current_event_index": 1,
        "events": [
            {"rounds": [{"status": "pending", "round_name": "Event 1 Heat"}]},
            {"rounds": [{"status": "pending", "round_name": "Event 2 Heat"}]},
        ],
    }

    event_index, _event, round_object = get_next_incomplete_round(tournament)

    assert event_index == 1
    assert round_object["round_name"] == "Event 2 Heat"


def test_multi_event_schedule_rejects_legacy_bracket_state(monkeypatch, capsys):
    event = {
        "event_name": "Bracket SB",
        "event_type": "bracket",
        "format": "bracket",
        "status": "ready",
        "rounds": [],
    }
    tournament = {"events": [event]}
    _script_input(monkeypatch, [""])

    returned = generate_complete_day_schedule(tournament)

    assert returned is tournament
    assert event["rounds"] == []
    assert "single-event tournament" in capsys.readouterr().out


def test_multi_event_results_reject_legacy_bracket_before_result_writes(monkeypatch, capsys):
    event = {
        "event_name": "Bracket SB",
        "event_type": "bracket",
        "format": "bracket",
        "status": "in_progress",
        "rounds": [
            {
                "round_name": "Heat 1",
                "round_type": "heat",
                "competitors": ["Alice", "Bob"],
                "status": "pending",
            }
        ],
    }
    tournament = {"events": [event]}
    monkeypatch.setattr(
        "woodchopping.ui.multi_event_ui.append_results_to_excel",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(AssertionError("canonical writer reached")),
    )
    _script_input(monkeypatch, [""])

    returned = sequential_results_workflow(tournament, {}, pd.DataFrame())

    assert returned is tournament
    assert event["rounds"][0]["status"] == "pending"
    assert "single-event tournament" in capsys.readouterr().out


def test_multi_event_heats_to_final_replays_through_save_boundaries(tmp_path, monkeypatch):
    names = ["Alice", "Bob", "Carol", "Drew"]
    roster = _competitors(names)
    heats = distribute_competitors_into_heats(roster, _handicaps(names), 2, 2)
    for heat in heats:
        heat["num_to_advance"] = 1
    event = {
        "event_name": "Replay Championship",
        "event_order": 1,
        "event_type": "championship",
        "event_code": "SB",
        "format": "heats_to_finals",
        "status": "in_progress",
        "wood_species": "S01",
        "wood_diameter": 300,
        "wood_quality": 5,
        "num_stands": 2,
        "all_competitors": names,
        "all_competitors_df": roster,
        "rounds": heats,
    }
    tournament = {
        "tournament_name": "Replay Day",
        "tournament_date": "2026-08-18",
        "events": [event],
        "total_events": 1,
        "events_completed": 0,
        "current_event_index": 0,
    }
    state_path = tmp_path / "multi_tournament_state.json"

    def record_results(_heat_df, _wood, round_object=None, **_kwargs):
        assert round_object is not None
        round_object["finish_order"] = {name: position for position, name in enumerate(round_object["competitors"], 1)}
        round_object["actual_results"] = {
            name: 29.0 + position for position, name in enumerate(round_object["competitors"], 1)
        }
        round_object["status"] = "in_progress"
        return True

    monkeypatch.setattr("woodchopping.ui.multi_event_ui.append_results_to_excel", record_results)
    monkeypatch.setattr(
        "woodchopping.ui.multi_event_ui.auto_save_multi_event",
        lambda state: save_multi_event_tournament(state, str(state_path)),
    )

    for _stage in range(2):
        _script_input(monkeypatch, ["1", "", "", "4"])
        sequential_results_workflow(tournament, {}, pd.DataFrame())
        tournament = load_multi_event_tournament(str(state_path))
        assert tournament is not None

    finals = [
        round_object for round_object in tournament["events"][0]["rounds"] if round_object["round_type"] == "final"
    ]
    assert len(finals) == 1

    _script_input(monkeypatch, ["1", "", ""])
    sequential_results_workflow(tournament, {}, pd.DataFrame())
    tournament = load_multi_event_tournament(str(state_path))

    assert tournament is not None
    assert tournament["events_completed"] == 1
    assert tournament["events"][0]["status"] == "completed"
    assert tournament["events"][0]["final_results"]["first_place"]


def test_multi_event_single_heat_summary_uses_completed_heat(monkeypatch, capsys):
    event = {
        "event_name": "300mm SB",
        "event_order": 1,
        "event_type": "championship",
        "event_code": "SB",
        "format": "single_heat",
        "status": "in_progress",
        "wood_species": "S01",
        "wood_diameter": 300,
        "wood_quality": 5,
        "all_competitors": ["Alice", "Bob"],
        "rounds": [
            {
                "round_name": "Heat 1",
                "round_type": "heat",
                "competitors": ["Alice", "Bob"],
                "status": "in_progress",
                "finish_order": {"Alice": 1, "Bob": 2},
                "actual_results": {"Alice": 30.0, "Bob": 32.0},
            }
        ],
    }
    tournament = {
        "tournament_name": "Replay Day",
        "tournament_date": "2026-08-18",
        "events": [event],
        "total_events": 1,
        "events_completed": 0,
    }
    complete_event_round(tournament, event, event["rounds"][0])
    _script_input(monkeypatch, [""])

    generate_tournament_summary(tournament)

    output = capsys.readouterr().out
    assert "1st Place: Alice (30.00s)" in output


def test_single_event_single_heat_final_summary_uses_completed_heat(monkeypatch, capsys):
    state = {
        "event_name": "Single Heat Championship",
        "format": "single_heat",
        "rounds": [
            {
                "round_name": "Heat 1",
                "round_type": "heat",
                "competitors": ["Alice", "Bob"],
                "status": "in_progress",
                "finish_order": {"Alice": 1, "Bob": 2},
                "actual_results": {"Alice": 30.0, "Bob": 32.0},
            }
        ],
    }
    assert complete_recorded_round(state, state["rounds"][0], entry_succeeded=True)
    _script_input(monkeypatch, [""])

    display_single_event_final_results(state)

    output = capsys.readouterr().out
    assert "1st" in output
    assert "Alice" in output


def test_bracket_with_byes_resumes_to_champion_without_result_writes(tmp_path):
    names = [f"Seed {seed}" for seed in range(1, 7)]
    predictions = {name: {"seed": seed} for seed, name in enumerate(names, 1)}
    state = initialize_bracket_tournament(num_stands=2, tentative_competitors=len(names))
    state["predictions"] = predictions
    state["all_competitors"] = names
    state["all_competitors_df"] = _competitors(names)
    state["rounds"] = generate_bracket_with_byes(predictions)
    state["total_rounds"] = len(state["rounds"])
    state["total_matches"] = sum(len(item["matches"]) for item in state["rounds"])

    first_next_round = state["rounds"][1]["matches"]
    populated = {
        competitor
        for match in first_next_round
        for competitor in (match["competitor1"], match["competitor2"])
        if competitor
    }
    assert {"Seed 1", "Seed 2"} <= populated

    save_path = tmp_path / "bracket_state.json"
    recorded = 0
    while (match := get_current_match(state)) is not None:
        record_match_result(state, match["match_id"], 29.0, 31.0, 1, 2)
        recorded += 1
        if recorded == 1:
            save_tournament_state(state, str(save_path))
            state = load_tournament_state(str(save_path))
            assert state is not None

    assert state["champion"] is not None
    assert state["runner_up"] is not None
    assert state["completed_matches"] == state["total_matches"]
    assert not list(tmp_path.glob("*.xlsx"))
    assert not list(tmp_path.glob("*.db"))
