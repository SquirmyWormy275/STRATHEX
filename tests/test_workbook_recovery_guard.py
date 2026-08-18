"""Regression coverage for fail-closed workbook result writes.

Every test redirects the production workbook path to ``tmp_path``. No test reads
or writes the repository's real competition workbook.
"""

from __future__ import annotations

import builtins
import hashlib
from pathlib import Path
from types import SimpleNamespace

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

import woodchopping.data as data_api
import woodchopping.data.excel_io as excel_io
from woodchopping.data.workbook_guard import guarded_append_results_to_excel


def _fake_paths(path: Path) -> SimpleNamespace:
    return SimpleNamespace(
        EXCEL_FILE=str(path),
        WOOD_SHEET="wood",
        COMPETITOR_SHEET="Competitor",
        RESULTS_SHEET="Results",
    )


def _write_complete_workbook(path: Path) -> None:
    workbook = Workbook()
    workbook.remove(workbook["Sheet"])

    wood = workbook.create_sheet("wood")
    wood.append(excel_io.WOOD_HEADERS)
    wood.append(
        [
            "Pinus strobus",
            "Eastern White Pine",
            "S01",
            "USA",
            "East",
            380,
            0.35,
            1,
            1,
            1,
            1,
        ]
    )

    competitors = workbook.create_sheet("Competitor")
    competitors.append(excel_io.COMPETITOR_HEADERS)
    competitors.append(["C001", "Alice Axe", "USA", "MT", "F"])

    results = workbook.create_sheet("Results")
    results.append(
        [
            "CompetitorID",
            "Event",
            "Time (seconds)",
            "Size (mm)",
            "Species Code",
            "Quality",
            "HeatID",
            "Date",
        ]
    )

    workbook.save(path)
    workbook.close()


def _digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _sheet_names(path: Path) -> set[str]:
    workbook = load_workbook(path, read_only=True)
    try:
        return {name.lower() for name in workbook.sheetnames}
    finally:
        workbook.close()


def test_unreadable_existing_workbook_is_not_touched(tmp_path, monkeypatch, capsys):
    path = tmp_path / "woodchopping.xlsx"
    original = b"this is not an xlsx file"
    path.write_bytes(original)
    monkeypatch.setattr(excel_io, "paths", _fake_paths(path))

    called = False

    def append_function(*_args, **_kwargs):
        nonlocal called
        called = True

    guarded_append_results_to_excel(append_function, excel_io)

    assert called is False
    assert path.read_bytes() == original
    output = capsys.readouterr().out
    assert "Results were not written" in output
    assert "left unchanged" in output


def test_guard_disables_results_only_workbook_fallback(tmp_path, monkeypatch):
    path = tmp_path / "woodchopping.xlsx"
    _write_complete_workbook(path)
    monkeypatch.setattr(excel_io, "paths", _fake_paths(path))
    before = _digest(path)
    original_factory = excel_io.Workbook

    def append_function():
        # This is the exact dangerous fallback used by the legacy append path.
        excel_io.Workbook()

    guarded_append_results_to_excel(append_function, excel_io)

    assert _digest(path) == before
    assert _sheet_names(path) == {"wood", "competitor", "results"}
    assert excel_io.Workbook is original_factory


def test_corrupt_post_write_file_restores_prewrite_backup(tmp_path, monkeypatch, capsys):
    path = tmp_path / "woodchopping.xlsx"
    _write_complete_workbook(path)
    monkeypatch.setattr(excel_io, "paths", _fake_paths(path))
    before = _digest(path)

    def corrupt_append():
        path.write_bytes(b"corrupt replacement")

    guarded_append_results_to_excel(corrupt_append, excel_io)

    assert _digest(path) == before
    assert _sheet_names(path) == {"wood", "competitor", "results"}
    assert "original workbook was restored" in capsys.readouterr().out


def test_valid_write_is_retained_and_temporary_backup_removed(tmp_path, monkeypatch):
    path = tmp_path / "woodchopping.xlsx"
    _write_complete_workbook(path)
    monkeypatch.setattr(excel_io, "paths", _fake_paths(path))

    def valid_append():
        workbook = load_workbook(path)
        workbook["Results"].append(["C001", "SB", 30.0, 300, "S01", 5, "H1", "2026-08-18"])
        workbook.save(path)
        workbook.close()
        return "saved"

    result = guarded_append_results_to_excel(valid_append, excel_io)

    assert result == "saved"
    workbook = load_workbook(path, read_only=True)
    try:
        assert workbook["Results"].max_row == 2
    finally:
        workbook.close()
    assert not list(tmp_path.glob("*.prewrite-backup"))
    assert not list(tmp_path.glob(".*.prewrite-backup"))


def test_missing_workbook_uses_complete_schema_guard(tmp_path, monkeypatch):
    path = tmp_path / "woodchopping.xlsx"
    monkeypatch.setattr(excel_io, "paths", _fake_paths(path))

    observed = {}

    def append_function():
        observed["sheets"] = _sheet_names(path)

    guarded_append_results_to_excel(append_function, excel_io)

    assert path.exists()
    assert observed["sheets"] == {"wood", "competitor", "results"}
    assert _sheet_names(path) == {"wood", "competitor", "results"}


def test_public_append_wrapper_blocks_legacy_partial_recovery(
    tmp_path,
    monkeypatch,
    capsys,
):
    path = tmp_path / "woodchopping.xlsx"
    _write_complete_workbook(path)
    monkeypatch.setattr(excel_io, "paths", _fake_paths(path))
    before = _digest(path)

    def fail_existing_workbook_load(*_args, **_kwargs):
        raise OSError("simulated transient workbook load failure")

    # The outer guard validates with its own read-only loader. This patch reaches
    # only the legacy routine's nested load/fallback branch.
    monkeypatch.setattr(excel_io, "load_workbook", fail_existing_workbook_load)

    answers = iter(["H1", "2", "30.0"])
    monkeypatch.setattr(builtins, "input", lambda _prompt="": next(answers))

    heat = pd.DataFrame(
        {
            "competitor_name": ["Alice Axe"],
            "mark": [3],
        }
    )
    wood = {
        "event": "SB",
        "species": "S01",
        "size_mm": 300,
        "quality": 5,
    }

    data_api.append_results_to_excel(heat, wood)

    assert _digest(path) == before
    assert _sheet_names(path) == {"wood", "competitor", "results"}
    assert "Refusing to create a replacement workbook" in capsys.readouterr().out


@pytest.mark.parametrize("missing_sheet", ["wood", "Competitor", "Results"])
def test_incomplete_existing_workbook_is_rejected(
    tmp_path,
    monkeypatch,
    missing_sheet,
):
    path = tmp_path / "woodchopping.xlsx"
    _write_complete_workbook(path)
    workbook = load_workbook(path)
    workbook.remove(workbook[missing_sheet])
    workbook.save(path)
    workbook.close()
    monkeypatch.setattr(excel_io, "paths", _fake_paths(path))
    before = _digest(path)

    called = False

    def append_function():
        nonlocal called
        called = True

    guarded_append_results_to_excel(append_function, excel_io)

    assert called is False
    assert _digest(path) == before
