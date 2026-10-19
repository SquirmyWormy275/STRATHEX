"""Adversarial regressions for the October 2026 Security Cloud findings."""

from __future__ import annotations

import ast
import copy
import io
import json
from pathlib import Path
from types import SimpleNamespace

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

from woodchopping.ui import multi_event_ui, personnel_ui, scratch_management
from woodchopping.ui.bracket_ui import generate_bracket_html_structure
from woodchopping.ui.terminal_rendering import TerminalOutput, sanitize_text
from woodchopping.ui.workflow_integrity import (
    approve_marks,
    bind_calculation,
    field_context_digest,
    new_event_id,
    prepare_field_change,
    require_approved_marks,
    require_removable_event,
)


def _state():
    state = {
        "event_id": "event_2",
        "event_name": "Synthetic event",
        "event_order": 2,
        "all_competitors": ["Alice", "Bob"],
        "format": "single_heat",
        "num_stands": 2,
        "all_competitors_df": pd.DataFrame({"competitor_name": ["Alice", "Bob"], "competitor_id": ["C1", "C2"]}),
        "wood_species": "Aspen",
        "wood_diameter": 300,
        "wood_quality": 7,
        "event_code": "SB",
        "prediction_as_of": "2026-10-18",
        "prediction_authority_ref": {"revision": 1},
        "handicap_results_all": [
            {"name": "Alice", "competitor_id": "C1", "mark": 3},
            {"name": "Bob", "competitor_id": "C2", "mark": 15},
        ],
        "rounds": [],
        "status": "ready",
    }
    bind_calculation(state)
    return state


@pytest.mark.parametrize(
    "text", ["=1+1", "+SUM(A1:A2)", "-1+2", "@SUM(A1:A2)", "\t=1+1", "\r=1+1", "  =1+1", "＝1+1", "José 山田"]
)
def test_roster_text_round_trips_as_literal_cells(tmp_path, monkeypatch, text):
    path = tmp_path / "roster.xlsx"
    workbook = Workbook()
    workbook.active.title = personnel_ui.COMPETITOR_SHEET
    workbook.active.append(["CompetitorID", "Name", "Country", "State/Province", "Gender"])
    workbook.save(path)
    workbook.close()
    monkeypatch.setattr(personnel_ui, "COMPETITOR_FILE", path)
    identity = personnel_ui._add_competitor_to_workbook(text, text, text, text)
    reloaded = load_workbook(path, data_only=False)
    try:
        cells = list(reloaded.active.iter_rows(min_row=2))[0]
        assert cells[0].value == identity
        assert all(cell.value == text.replace("\r", "\n") and cell.data_type == "s" for cell in cells[1:])
    finally:
        reloaded.close()


def test_bracket_escapes_all_untrusted_html_contexts():
    payload = '<img src=x onerror="alert(1)"> & </title><script>alert(2)</script>'
    match = dict(
        status='x" onclick="alert(3)',
        competitor1=payload,
        competitor2=None,
        seed1=payload,
        seed2=None,
        time1=None,
        time2=None,
        finish_position1=1,
        finish_position2=None,
        winner=payload,
        advances_to=payload,
        match_id=payload,
    )
    state = dict(
        event_name=payload,
        wood_species=payload,
        wood_diameter=payload,
        wood_quality=payload,
        num_competitors=payload,
        total_rounds=payload,
        completed_matches=payload,
        total_matches=payload,
        rounds=[dict(round_name=payload, matches=[match])],
    )
    before = copy.deepcopy(state)
    html = generate_bracket_html_structure(state)
    assert "<script>" not in html and "<img" not in html and 'onclick="' not in html
    assert "&lt;img" in html and "&amp;" in html
    assert "Content-Security-Policy" in html
    assert state == before


@pytest.mark.parametrize(
    "control",
    [
        "\x1b[2J",
        "\x1b]52;c;YXR0YWNr\x07",
        "\x1b]8;;https://invalid\x1b\\",
        "\x1bPattack\x1b\\",
        "\x9d52;c;x\x9c",
        "\x1b_attack\x1b\\",
        "\x07\r\b",
        "\u202e\u200b",
    ],
)
def test_terminal_controls_never_reach_the_output_stream(control):
    target = io.StringIO()
    terminal = TerminalOutput(target)
    for text in ["Roster: Alice", control, " Bob\n", "Result: ", control]:
        terminal.write(text)
        terminal.flush()
    output = target.getvalue()
    assert all(ord(c) >= 32 or c == "\n" for c in output)
    assert not any(127 <= ord(c) <= 159 for c in output)
    assert "\u202e" not in output and "\u200b" not in output
    assert sanitize_text("José 山田") == "José 山田"


def test_unapproved_sheet_cannot_be_scheduled():
    state = _state()
    with pytest.raises(ValueError, match="approve"):
        require_approved_marks(state)
    approve_marks(state, "JD", "2026-10-18T12:00:00Z")
    require_approved_marks(state)


@pytest.mark.parametrize(
    "change",
    [
        "wood_species",
        "wood_diameter",
        "wood_quality",
        "event_code",
        "prediction_as_of",
        "prediction_authority_ref",
        "all_competitors",
        "mark",
        "competitor_id",
    ],
)
def test_approval_invalidates_on_any_dependent_change(change):
    state = _state()
    approve_marks(state, "JD", "2026-10-18T12:00:00Z")
    if change == "mark":
        state["handicap_results_all"][0]["mark"] = 50
    elif change == "competitor_id":
        state["all_competitors_df"].loc[0, "competitor_id"] = "SWAPPED"
    elif change == "all_competitors":
        state[change].reverse()
    else:
        state[change] = "changed"
    with pytest.raises(ValueError):
        require_approved_marks(state)


def test_round_marks_must_match_reviewed_sheet():
    state = _state()
    approve_marks(state, "JD", "2026-10-18T12:00:00Z")
    state["rounds"] = [
        {
            "competitors": ["Alice", "Bob"],
            "status": "pending",
            "handicap_results": copy.deepcopy(state["handicap_results_all"]),
        }
    ]
    state["rounds"][0]["handicap_results"][1]["mark"] = 99
    with pytest.raises(ValueError, match="Scheduled marks"):
        require_approved_marks(state)


def test_approved_context_survives_json_roundtrip():
    state = _state()
    approve_marks(state, "JD", "2026-10-18T12:00:00Z")
    frame = state.pop("all_competitors_df")
    loaded = json.loads(json.dumps(state))
    loaded["all_competitors_df"] = frame.copy()
    require_approved_marks(loaded)


def test_changed_draft_archives_marks_and_clears_schedule():
    state = _state()
    approve_marks(state, "JD", "now")
    state["rounds"] = [{"status": "pending", "competitors": ["Alice", "Bob"]}]
    root = {"events": [state], "schedule": ["old field"]}
    original = copy.deepcopy(state["handicap_results_all"])
    assert prepare_field_change(root, reason="scratch")
    assert state["handicap_results_all"] == [] and state["rounds"] == []
    assert "mark_approval" not in state and root["schedule"] == []
    assert state["superseded_drafts"][0]["artifacts"]["handicap_results_all"] == original


@pytest.mark.parametrize(
    "evidence",
    [{"status": "completed"}, {"actual_results": {"Alice": 10}}, {"status": "in_progress"}, {"issued": True}],
)
def test_field_change_preserves_official_results_atomically(evidence):
    first, second = _state(), _state()
    second["rounds"] = [evidence]
    before = field_context_digest(first), copy.deepcopy(second["rounds"])
    root = {"events": [first, second], "schedule": ["keep"]}
    assert not prepare_field_change(root, reason="assignment")
    assert first["handicap_results_all"] and root["schedule"] == ["keep"]
    assert before == (field_context_digest(first), second["rounds"])


def test_v3_evidence_prevents_draft_reset():
    state = _state()
    state["v3_issue_receipt"] = {"issued": True}
    assert not prepare_field_change(state, reason="wood")
    assert state["v3_issue_receipt"] and state["handicap_results_all"]


@pytest.mark.parametrize("reference", ["assignments", "fees", "schedule", "results", "calculation"])
def test_referenced_events_cannot_be_deleted(reference):
    state = _state()
    state["all_competitors"] = []
    state["handicap_results_all"] = []
    root = {"events": [state]}
    if reference == "assignments":
        root["tournament_roster"] = [{"events_entered": [state["event_id"]]}]
    elif reference == "fees":
        root["tournament_roster"] = [{"entry_fees_paid": {state["event_id"]: False}}]
    elif reference == "schedule":
        root["schedule"] = ["field:event_2:heat-1"]
    elif reference == "results":
        state["rounds"] = [{"status": "completed", "actual_results": {"Alice": 10}}]
    else:
        state["superseded_drafts"] = [{"marks": [3]}]
    with pytest.raises(ValueError):
        require_removable_event(root, state)


def test_deleting_unused_draft_does_not_renumber_ids(monkeypatch):
    events = [
        {"event_id": f"event_{i}", "event_order": i, "event_name": f"Event {i}", "status": "pending", "rounds": []}
        for i in range(1, 4)
    ]
    root = {"events": events, "total_events": 3}
    answers = iter(["2", "DELETE", ""])
    monkeypatch.setattr("builtins.input", lambda _="": next(answers))
    monkeypatch.setattr(multi_event_ui, "auto_save_multi_event", lambda _: True)
    multi_event_ui.remove_event_from_tournament(root)
    assert [e["event_id"] for e in root["events"]] == ["event_1", "event_3"]
    assert [e["event_order"] for e in root["events"]] == [1, 2]
    assert new_event_id() not in {"event_1", "event_2", "event_3"}


def test_single_event_menu_blocks_direct_heat_generation_without_approval():
    # Execute only the function: importing MainProgram opens the operator workbook.
    path = Path(__file__).resolve().parents[1] / "MainProgramV5_2.py"
    function = next(
        node
        for node in ast.parse(path.read_text()).body
        if isinstance(node, ast.FunctionDef) and node.name == "single_event_menu"
    )
    output = []
    answers = iter(["9", "", "16"])

    def select(state, **_):
        state.update(_state())

    namespace = {
        "pd": pd,
        "os": SimpleNamespace(name="posix", system=lambda _: None),
        "input": lambda _="": next(answers),
        "print": lambda *args, **_: output.append(" ".join(map(str, args))),
        "select_prediction_engine_for_scope": select,
        "_prediction_authority_store": object(),
        "_v3_readiness_provider": None,
        "_judge_actor": lambda: "JD",
        "display_prediction_engine_banner": lambda *_: None,
        "get_wood_status_display": lambda _: "wood",
        "get_competitor_status_display": lambda _: "competitors",
        "resolve_prediction_engine": lambda *_: SimpleNamespace(engine="v2"),
        "require_approved_marks": require_approved_marks,
    }
    exec(compile(ast.Module(body=[function], type_ignores=[]), str(path), "exec"), namespace)
    namespace["single_event_menu"]()
    assert any("approve" in line for line in output)
    assert namespace["tournament_state"]["rounds"] == []
