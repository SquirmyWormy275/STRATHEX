"""Synthetic official-use, workbook and input-boundary regressions."""

import io
import zipfile
from copy import deepcopy
from pathlib import Path
from types import SimpleNamespace

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

from woodchopping.data import excel_io, workbook_access
from woodchopping.ui import state_persistence
from woodchopping.ui.official_eligibility import require_official_eligibility
from woodchopping.ui.schedule_printout import _generate_single_event_schedule, generate_printable_schedule
from woodchopping.ui.workflow_integrity import approve_marks, bind_calculation


def _workbook(path, identity="C1"):
    wb = Workbook()
    wb.active.title = "wood"
    wb.active.append(excel_io.WOOD_HEADERS)
    sheet = wb.create_sheet("Competitor")
    sheet.append(excel_io.COMPETITOR_HEADERS)
    sheet.append([identity, "Alice", "USA", "MT", "F"])
    excel_io.detect_results_sheet(wb)
    wb.save(path)
    wb.close()


@pytest.fixture
def official_field(tmp_path, monkeypatch):
    path = tmp_path / "synthetic.xlsx"
    monkeypatch.setattr(
        excel_io,
        "paths",
        SimpleNamespace(
            EXCEL_FILE=str(path), WOOD_SHEET="wood", COMPETITOR_SHEET="Competitor", RESULTS_SHEET="Results"
        ),
    )
    _workbook(path)
    rows = [{"name": "Alice", "competitor_id": "C1", "mark": 3}]
    field = {
        "event_name": "Synthetic",
        "all_competitors": ["Alice"],
        "all_competitors_df": pd.DataFrame({"competitor_name": ["Alice"], "competitor_id": ["C1"]}),
        "handicap_results_all": rows,
        "rounds": [],
    }
    bind_calculation(field)
    approve_marks(field, "JD", "2026-10-19T00:00:00Z")
    round_object = {"competitors": ["Alice"], "round_name": "Final", "handicap_results": deepcopy(rows)}
    return path, field, round_object


def test_reviewed_result_rejects_live_id_rebinding_before_interaction(official_field, monkeypatch):
    path, field, round_object = official_field
    _workbook(path, "ATTACKER")
    before = path.read_bytes()
    monkeypatch.setattr("builtins.input", lambda *_: pytest.fail("rebound roster reached input"))
    monkeypatch.setattr("woodchopping.data.store_registry.get_store", lambda: pytest.fail("ResultStore changed"))
    assert not excel_io.append_results_to_excel(
        None, {"event": "UH", "species": "Aspen", "size_mm": 300, "quality": 7}, round_object, field
    )
    assert path.read_bytes() == before
    assert "actual_results" not in round_object


def test_roster_change_during_entry_rejected_before_either_canonical_sink(official_field, monkeypatch):
    path, field, round_object = official_field
    answers = iter(["2", "20.5"])

    def answer(_prompt):
        value = next(answers)
        if value == "20.5":
            _workbook(path, "ATTACKER")
        return value

    monkeypatch.setattr("builtins.input", answer)
    monkeypatch.setattr("woodchopping.data.store_registry.get_store", lambda: pytest.fail("ResultStore changed"))
    from woodchopping.data import append_results_to_excel

    assert not append_results_to_excel(
        None, {"event": "UH", "species": "Aspen", "size_mm": 300, "quality": 7}, round_object, field
    )
    wb = load_workbook(path)
    assert wb["Results"].max_row == 1
    wb.close()
    assert "actual_results" not in round_object


def test_official_schedule_encodes_controls_without_altering_identity():
    name = "Alice\nStand 99\u202e\u200b"
    state = {
        "event_name": "Show\r\nFORGED",
        "rounds": [
            {"round_name": "Final\u2028FORGED", "competitors": [name], "handicap_results": [{"name": name, "mark": 17}]}
        ],
    }
    original = deepcopy(state)
    text = _generate_single_event_schedule(state)
    assert state == original
    assert text.count("Stand 1:") == 1
    assert "(Mark 17)" in text
    assert not any(c in text for c in ("\r", "\u202e", "\u200b", "\u2028"))
    assert "\\u000a" in text
    assert all(line.startswith(("║", "╔", "╠", "╚")) for line in text.splitlines())


@pytest.mark.parametrize("issued", [False, True])
def test_rehearsal_cannot_export_or_write_official_results(tmp_path, monkeypatch, issued):
    from test_v3_issue_settlement_workflow import _field_rows, _v3_state

    from woodchopping.ui.multi_event_ui import record_and_settle_v3_round

    state, store = _v3_state(tmp_path)
    rows = _field_rows()
    field = {"competitors": ["Alice", "Bob"], "handicap_results": rows}
    state["rounds"] = [field]
    if issued:
        state["v3_issue_batches"] = {"receipt:heat-one": "issue_batch:one"}
    monkeypatch.chdir(tmp_path)
    with pytest.raises(ValueError, match="rehearsal"):
        require_official_eligibility(state, authority_store=store)
    with pytest.raises(ValueError, match="rehearsal"):
        generate_printable_schedule(state, authority_store=store)
    assert not record_and_settle_v3_round(
        state,
        {},
        field,
        write_action=lambda: pytest.fail("canonical write reached"),
        authority_store=store,
        v3_adapter=object(),
    )
    assert not field.get("canonical_results_recorded")
    assert not list(tmp_path.glob("*schedule*.txt"))
    from woodchopping.ui import handicap_ui, multi_event_ui

    monkeypatch.setattr(multi_event_ui, "_prediction_authority_store", store)
    monkeypatch.setattr("builtins.input", lambda *_: pytest.fail("rehearsal reached legacy result input"))
    monkeypatch.setattr("woodchopping.data.store_registry.get_store", lambda: pytest.fail("ResultStore changed"))
    for writer in (excel_io.append_results_to_excel, handicap_ui.append_results_to_excel):
        assert not writer(None, {"event": "UH"}, field, state)


def test_local_unissued_results_block_before_collection(tmp_path):
    from test_v3_issue_settlement_workflow import _field_rows, _v3_state

    from woodchopping.ui.multi_event_ui import record_and_settle_v3_round

    state, store = _v3_state(tmp_path, local=True)
    field = {"handicap_results": _field_rows()}
    assert not record_and_settle_v3_round(
        state,
        {},
        field,
        write_action=lambda: pytest.fail("canonical write reached"),
        authority_store=store,
        v3_adapter=SimpleNamespace(requires_explicit_issue=True),
        input_fn=lambda *_: pytest.fail("unissued field reached input"),
    )


def test_unselected_durable_scope_cannot_authorize_official_use(tmp_path):
    from woodchopping.ui.prediction_context import PredictionAuthorityStore, attach_authority_reference

    store = PredictionAuthorityStore(tmp_path / "authority.sqlite3")
    scope = store.create_scope(owner_kind="tournament")
    state = {}
    attach_authority_reference(state, scope.reference)
    with pytest.raises(ValueError, match="Select and review"):
        require_official_eligibility(state, authority_store=store)


def test_oversized_untrusted_save_rejected_before_read_or_json(tmp_path, monkeypatch):
    path = tmp_path / "large.json"
    with path.open("wb") as stream:
        stream.truncate(state_persistence.MAX_SAVE_BYTES + 1)
    monkeypatch.setattr(state_persistence.json, "loads", lambda *_: pytest.fail("untrusted oversized JSON parsed"))
    with pytest.raises(ValueError, match="budget"):
        state_persistence._load_with_recovery(
            str(path), lambda *_: None, authority_store=SimpleNamespace(save_payload_digests=lambda _: ("a" * 64, None))
        )


def test_workbook_dimension_budget_before_openpyxl(official_field, monkeypatch):
    path, _, _ = official_field
    raw = path.read_bytes()
    with zipfile.ZipFile(io.BytesIO(raw)) as source, zipfile.ZipFile(path, "w") as target:
        for entry in source.infolist():
            payload = source.read(entry)
            if entry.filename == "xl/worksheets/sheet1.xml":
                payload = payload.replace(b'ref="A1:K1"', b'ref="A1:XFD1048576"')
            target.writestr(entry, payload)
    monkeypatch.setattr(
        workbook_access, "_load_workbook", lambda *_args, **_kwargs: pytest.fail("oversized workbook parsed")
    )
    with pytest.raises(ValueError, match="dimension"):
        workbook_access.load_workbook(path)


def test_changed_workbook_cannot_redirect_save_to_victim(official_field, tmp_path):
    path, _, _ = official_field
    workbook = workbook_access.load_workbook(path)
    victim = tmp_path / "victim"
    victim.write_bytes(b"retained evidence")
    path.unlink()
    path.symlink_to(victim)
    with pytest.raises((ValueError, PermissionError)):
        workbook_access.save_workbook(workbook, path)
    workbook.close()
    assert victim.read_bytes() == b"retained evidence"


def test_concurrent_regular_workbook_is_preserved(official_field):
    path, _, _ = official_field
    workbook = workbook_access.load_workbook(path)
    _workbook(path, "NEW")
    current = path.read_bytes()
    with pytest.raises(RuntimeError, match="changed"):
        workbook_access.save_workbook(workbook, path)
    workbook.close()
    assert path.read_bytes() == current
