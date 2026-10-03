"""Recovered UI defects, exercised with synthetic state and disposable workbooks."""

from __future__ import annotations

import ast
from copy import deepcopy
from dataclasses import replace
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

from woodchopping.ui import history_entry, personnel_ui, state_persistence, v52_helpers
from woodchopping.ui.bracket_ui import render_match_box_compact


def test_id_only_history_loads_all_rows_with_roster_display_names(tmp_path, monkeypatch, capsys):
    from woodchopping.data import excel_io

    path = tmp_path / "history.xlsx"
    workbook = Workbook()
    roster = workbook.active
    roster.title = "Competitor"
    roster.append(["CompetitorID", "Name", "Country", "State/Province", "Gender"])
    roster.append(["SYN001", "Same Name", "SYN", "", "M"])
    roster.append(["SYN002", "Same Name", "SYN", "", "M"])
    history = workbook.create_sheet("Results")
    history.append(["CompetitorID", "Event", "Time (seconds)", "Size (mm)", "Species Code", "Date"])
    history.append(["SYN001", "uh", 30, 300, "S01", "2024-01-01"])
    history.append(["SYN001", "UH", 31, 300, "S01", "2025-01-01"])
    history.append(["SYN002", "SB", 40, 275, "S01", "2025-01-02"])
    workbook.save(path)
    workbook.close()
    before = path.read_bytes()
    monkeypatch.setattr(excel_io, "paths", replace(excel_io.paths, EXCEL_FILE=str(path)))

    results = excel_io.load_results_df()

    assert results["competitor_id"].tolist() == ["SYN001", "SYN001", "SYN002"]
    assert results["competitor_name"].tolist() == ["Same Name [SYN001]", "Same Name [SYN001]", "Same Name [SYN002]"]
    assert results["raw_time"].tolist() == [30, 31, 40]
    assert results["event"].tolist() == ["UH", "UH", "SB"]
    assert results["date"].notna().all()
    assert path.read_bytes() == before
    assert "Error loading results" not in capsys.readouterr().out


@pytest.mark.parametrize("decision,created", [("1", 1), ("", 0)])
def test_finance_before_tournament_handles_create_and_cancel(decision, created):
    # Importing the legacy entry script initializes operator data. Execute its
    # actual menu function without that initialization or any real workbook.
    path = Path(__file__).resolve().parents[1] / "MainProgramV5_2.py"
    function = next(
        node
        for node in ast.parse(path.read_text(encoding="utf-8")).body
        if isinstance(node, ast.FunctionDef) and node.name == "multi_event_tournament_menu"
    )
    calls = []
    answers = iter(["5"])
    namespace = {
        "multi_event_tournament_state": {},
        "comp_df": pd.DataFrame(),
        "input": lambda _prompt="": next(answers),
        "print": lambda *args, **kwargs: None,
        "display_actionable_error": lambda *args, **kwargs: decision,
        "_create_multi_event_with_engine": lambda: calls.append(True) or {},
    }
    exec(compile(ast.Module(body=[function], type_ignores=[]), str(path), "exec"), namespace)
    with pytest.raises(StopIteration):
        namespace["multi_event_tournament_menu"]()
    assert len(calls) == created


def test_unfilled_double_elimination_match_renders_tbd(capsys):
    render_match_box_compact({"match_id": "LR2-M1", "status": "pending", "competitor1": None, "competitor2": None})
    assert "TBD" in capsys.readouterr().out


def _roster(tmp_path, monkeypatch, ids):
    path = tmp_path / "roster.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = personnel_ui.COMPETITOR_SHEET
    sheet.append(["CompetitorID", "Name", "Country", "State/Province", "Gender"])
    for index, identity in enumerate(ids):
        sheet.append([identity, f"Existing {index}", "USA", "", "M"])
    workbook.save(path)
    workbook.close()
    monkeypatch.setattr(personnel_ui, "COMPETITOR_FILE", str(path))
    monkeypatch.setattr(personnel_ui, "ensure_workbook", lambda *args: None)
    monkeypatch.setattr(personnel_ui, "load_competitors_df", lambda: pd.DataFrame())
    monkeypatch.setattr(personnel_ui, "add_historical_times_for_competitor", lambda *args: None)
    answers = iter(["New Competitor", "USA", "", "M"])
    monkeypatch.setattr("builtins.input", lambda _prompt="": next(answers))
    return path


def test_new_competitor_does_not_reuse_deleted_roster_row_id(tmp_path, monkeypatch):
    path = _roster(tmp_path, monkeypatch, ["C001", "C003"])
    personnel_ui.add_competitor_with_times()
    workbook = load_workbook(path, read_only=True)
    ids = [row[0] for row in workbook[personnel_ui.COMPETITOR_SHEET].iter_rows(min_row=2, values_only=True)]
    workbook.close()
    assert ids[:2] == ["C001", "C003"]
    assert len(ids) == len(set(ids)) == 3


def test_duplicate_existing_competitor_ids_block_roster_write(tmp_path, monkeypatch, capsys):
    path = _roster(tmp_path, monkeypatch, ["C003", "C003"])
    before = path.read_bytes()
    personnel_ui.add_competitor_with_times()
    assert path.read_bytes() == before
    assert "duplicate" in capsys.readouterr().out.lower()


@pytest.mark.parametrize("selection", ["0", "-1"])
def test_entry_editor_rejects_nonpositive_competitor_selection(monkeypatch, selection, capsys):
    state = {
        "events": [
            {
                "event_id": "event-1",
                "event_name": "Underhand",
                "all_competitors": [],
                "rounds": [],
                "competitor_status": {},
            }
        ],
        "tournament_roster": [{"competitor_name": "Alice", "events_entered": [], "entry_fees_paid": {}}],
        "entry_fee_tracking_enabled": False,
    }
    original = deepcopy(state)
    # If invalid input is accepted, the old code consumes "4" as an event
    # selection instead of returning directly to the editor menu.
    answers = iter(["1", selection, "4"])
    monkeypatch.setattr("builtins.input", lambda _prompt="": next(answers))
    v52_helpers.edit_event_entries(state)
    assert state == original
    assert "Invalid selection" in capsys.readouterr().out


@pytest.mark.parametrize("multi", [False, True])
def test_payout_positions_survive_validated_save_and_resume(tmp_path, multi):
    payout = {"enabled": True, "payouts": {1: 100.0, 2: 50.0}, "num_places": 2, "total_purse": 150.0}
    event = {"event_name": "Underhand", "rounds": [], "payout_config": payout}
    state = {"events": [event]} if multi else event
    save = state_persistence.save_multi_event_tournament if multi else state_persistence.save_tournament_state
    load = state_persistence.load_multi_event_tournament if multi else state_persistence.load_tournament_state
    path = str(tmp_path / "tournament.json")
    assert save(state, path)
    loaded = load(path)
    actual = loaded["events"][0] if multi else loaded
    assert actual["payout_config"]["payouts"][1] == 100.0
    assert payout["payouts"] == {1: 100.0, 2: 50.0}


def test_invalid_history_date_is_corrected_before_persisting(monkeypatch):
    answers = iter(["30", "not-a-date", "2025-02-30", "2025-02-10", ""])
    monkeypatch.setattr("builtins.input", lambda _prompt="": next(answers))
    writes = []
    monkeypatch.setattr(history_entry, "save_time_to_results", lambda **kwargs: writes.append(kwargs))
    assert history_entry.prompt_add_competitor_times("Alice", "UH", {"species": "Pine", "size_mm": 300})
    assert len(writes) == 1
    assert writes[0]["timestamp"] == "2025-02-10"


def test_wood_quality_rejects_out_of_range_instead_of_changing_it(monkeypatch, capsys):
    from woodchopping.ui.wood_ui import enter_wood_quality

    answers = iter(["0", "11", "5"])
    monkeypatch.setattr("builtins.input", lambda _prompt="": next(answers))
    selected = enter_wood_quality({"quality": 7})
    assert selected["quality"] == 5
    assert capsys.readouterr().out.count("Invalid input") == 2


def test_tournament_menu_wood_save_and_return_match_displayed_options():
    path = Path(__file__).resolve().parents[1] / "MainProgramV5_2.py"
    function = next(
        node
        for node in ast.parse(path.read_text(encoding="utf-8")).body
        if isinstance(node, ast.FunctionDef) and node.name == "multi_event_tournament_menu"
    )
    answers = iter(["16", "17", "", "18"])
    calls = []
    state = {"tournament_name": "Synthetic tournament"}
    namespace = {
        "multi_event_tournament_state": state,
        "comp_df": pd.DataFrame(),
        "input": lambda _prompt="": next(answers),
        "print": lambda *args, **kwargs: None,
        "view_wood_count": lambda value: calls.append(("wood", value)),
        "save_multi_event_tournament": lambda value, filename, **kwargs: calls.append(("save", filename)),
        "_prediction_authority_store": None,
        "display_tournament_progress_tracker": lambda *args: None,
        "display_prediction_engine_banner": lambda *args: None,
    }
    exec(compile(ast.Module(body=[function], type_ignores=[]), str(path), "exec"), namespace)
    namespace["multi_event_tournament_menu"]()
    assert calls == [("wood", state), ("save", "saves/multi_tournament_state.json")]


def test_personal_best_watch_uses_canonical_result_columns(monkeypatch, capsys):
    from woodchopping import data
    from woodchopping.ui.championship_simulator import _display_personal_best_watch

    history = pd.DataFrame([{"competitor_name": "Alice", "event": "UH", "raw_time": 30.0}])
    monkeypatch.setattr(data, "load_results_df", lambda: history)
    _display_personal_best_watch([{"name": "Alice", "predicted_time": 30.1}], {"event": "UH"})
    assert "Alice" in capsys.readouterr().out
