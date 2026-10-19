"""Real synthetic workbook coverage for Linux result publication boundaries."""

from __future__ import annotations

import copy
import os
import stat
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import Workbook, load_workbook

import config
from woodchopping.data import excel_io
from woodchopping.ui import linux_results


@pytest.fixture
def result_workbook(tmp_path, monkeypatch):
    path = tmp_path / "synthetic.xlsx"
    paths = SimpleNamespace(
        EXCEL_FILE=str(path), WOOD_SHEET="Wood", COMPETITOR_SHEET="Competitor", RESULTS_SHEET="Results"
    )
    monkeypatch.setattr(config, "paths", paths)
    monkeypatch.setattr(excel_io, "paths", paths)
    monkeypatch.setenv("STRATHEX_TEST_DB", "1")
    backup = tmp_path / "backups"
    backup.mkdir()
    monkeypatch.setenv("STRATHEX_WORKBOOK_BACKUP_DIR", str(backup))
    book = Workbook()
    book.active.title = "Competitor"
    book.active.append(excel_io.COMPETITOR_HEADERS)
    excel_io.append_literal_row(book.active, ["=1+1", "Synthetic Athlete", "US", "MT", "M"])
    book.create_sheet("Wood").append(excel_io.WOOD_HEADERS)
    book.create_sheet("Results").append(
        ["CompetitorID", "Event", "Time (seconds)", "Size (mm)", "Species Code", "Date", "Round", "HeatID"]
    )
    book.save(path)
    book.close()
    return path, paths


def _round(identifier="=1+1", settlement="synthetic-settlement"):
    return {
        "handicap_results": [
            {
                "local_competitor_id": identifier,
                "receipt_id": "synthetic-receipt",
                "name": "Synthetic Athlete",
                "mark": 3,
            }
        ],
        "official_outcomes": {
            identifier: {"status": "completion", "raw_time_ms": 25000, "penalty_ms": None, "official_placing": 1}
        },
        "v3_settlement_id": settlement,
        "round_name": "Synthetic heat",
    }


def _event():
    return {"event_code": "UH", "wood_diameter": 300, "wood_species": "Aspen", "wood_quality": 7}


def _rows(path):
    book = load_workbook(path, data_only=False)
    try:
        sheet = book["Results"]
        headers = [cell.value for cell in sheet[1]]
        return [
            {header: (cell.value, cell.data_type) for header, cell in zip(headers, row)} for row in list(sheet.rows)[1:]
        ]
    finally:
        book.close()


@pytest.mark.parametrize("identifier", ["=1+1", "+1+1", "-1+1", "@SUM(1,1)", "ordinary-id"])
def test_linux_export_preserves_literal_identifiers_and_numeric_results(result_workbook, identifier):
    path, _ = result_workbook
    assert excel_io.load_competitors_df()["competitor_id"].tolist() == ["=1+1"]
    round_object = _round(identifier)
    round_object["round_name"] = "=1+1"
    round_object["handicap_results"][0]["receipt_id"] = "=2+2"
    round_object["v3_settlement_id"] = "=3+3"
    event = {**_event(), "wood_species": "=4+4"}
    before = copy.deepcopy(round_object)
    assert linux_results.export_results(round_object, event, {}) is True
    assert round_object == before
    row = _rows(path)[0]
    for header, value in {
        "CompetitorID": identifier,
        "HeatID": "=1+1",
        "V3 Receipt": "=2+2",
        "V3 Settlement": "=3+3",
        "Species Code": "=4+4",
    }.items():
        assert row[header] == (value, "s")
    assert row["Time (seconds)"] == (25, "n")
    assert row["Official Placing"] == (1, "n")
    assert row["Size (mm)"] == (300, "n")
    assert not list(path.parent.glob(f".{path.name}.*.tmp"))


def test_linux_export_duplicate_and_correction_preserve_history(result_workbook):
    path, _ = result_workbook
    original = _round()
    linux_results.export_results(original, _event(), {})
    linux_results.export_results(original, _event(), {})
    assert len(_rows(path)) == 1
    revised = _round(settlement="synthetic-correction")
    revised["official_outcomes"]["=1+1"]["raw_time_ms"] = 26000
    revised["result_source_revision"] = 2
    linux_results.export_results(revised, _event(), {}, supersedes_settlement_id="synthetic-settlement")
    prior, current = _rows(path)
    assert prior["CompetitorID"] == current["CompetitorID"] == ("=1+1", "s")
    assert prior["Time (seconds)"] == ("SUPERSEDED", "s")
    assert prior["Original Raw Time (seconds)"] == (25, "n")
    assert current["Time (seconds)"] == (26, "n")
    assert current["Result Revision"] == (2, "n")


@pytest.mark.skipif(os.name != "posix", reason="POSIX sibling symlink attack")
def test_old_predictable_symlink_is_not_followed_or_published(result_workbook):
    path, _ = result_workbook
    victim = path.parent / "unrelated.txt"
    victim.write_bytes(b"keep this unrelated file")
    hostile = path.with_name(path.name + ".v3-next.xlsx")
    hostile.symlink_to(victim)
    linux_results.export_results(_round(), _event(), {})
    assert victim.read_bytes() == b"keep this unrelated file"
    assert hostile.is_symlink()
    assert not path.is_symlink()
    assert _rows(path)[0]["CompetitorID"] == ("=1+1", "s")
    assert stat.S_IMODE(path.stat().st_mode) == 0o600


@pytest.mark.skipif(os.name != "posix", reason="POSIX publication identity race")
def test_replaced_staging_entry_fails_closed(result_workbook, monkeypatch):
    path, _ = result_workbook
    before = path.read_bytes()
    victim = path.parent / "unrelated.txt"
    victim.write_bytes(b"keep this unrelated file")
    original_save = Workbook.save

    def swap_after_write(book, stream):
        assert hasattr(stream, "fileno")
        original_save(book, stream)
        (staged,) = path.parent.glob(f".{path.name}.*.tmp")
        staged.unlink()
        staged.symlink_to(victim)

    monkeypatch.setattr(Workbook, "save", swap_after_write)
    with pytest.raises(RuntimeError, match="staging file changed"):
        linux_results.export_results(_round(), _event(), {})
    assert path.read_bytes() == before and not path.is_symlink()
    assert victim.read_bytes() == b"keep this unrelated file"
    assert not list(path.parent.glob(f".{path.name}.*.tmp"))


def test_failed_save_preserves_original_and_removes_staging(result_workbook, monkeypatch):
    path, _ = result_workbook
    before = path.read_bytes()

    def fail_save(book, stream):
        stream.write(b"incomplete workbook")
        raise OSError("synthetic disk failure")

    monkeypatch.setattr(Workbook, "save", fail_save)
    with pytest.raises(OSError, match="synthetic disk failure"):
        linux_results.export_results(_round(), _event(), {})
    assert path.read_bytes() == before
    assert not list(path.parent.glob(f".{path.name}.*.tmp"))


def test_concurrent_workbook_change_is_not_overwritten(result_workbook, monkeypatch):
    path, _ = result_workbook
    original_save = Workbook.save

    def concurrent_write(book, stream):
        original_save(book, stream)
        path.write_bytes(b"another writer's updated workbook")

    monkeypatch.setattr(Workbook, "save", concurrent_write)
    with pytest.raises(RuntimeError, match="workbook changed during export"):
        linux_results.export_results(_round(), _event(), {})
    assert path.read_bytes() == b"another writer's updated workbook"
    assert not list(path.parent.glob(f".{path.name}.*.tmp"))


@pytest.mark.skipif(os.name != "posix", reason="POSIX owner-controlled parent modes")
def test_owner_controlled_publicly_readable_parent_is_not_tightened(result_workbook):
    path, _ = result_workbook
    path.parent.chmod(0o755)
    linux_results.export_results(_round(), _event(), {})
    assert stat.S_IMODE(path.parent.stat().st_mode) == 0o755
    assert _rows(path)[0]["CompetitorID"] == ("=1+1", "s")


@pytest.mark.skipif(os.name != "posix", reason="POSIX shared parent modes")
@pytest.mark.parametrize("mode", [0o777, 0o1777])
def test_shared_workbook_parent_rejected_without_tightening(result_workbook, mode):
    path, paths = result_workbook
    shared = path.parent / "shared"
    shared.mkdir()
    shared.chmod(mode)
    shared_path = shared / path.name
    path.rename(shared_path)
    paths.EXCEL_FILE = str(shared_path)
    before = shared_path.read_bytes()
    with pytest.raises(PermissionError, match="replaceable|sibling"):
        linux_results.export_results(_round(), _event(), {})
    assert shared_path.read_bytes() == before
    assert stat.S_IMODE(shared.stat().st_mode) == mode
    assert not list((path.parent / "backups").iterdir())


def test_legacy_result_entry_preserves_formula_shaped_roster_id(result_workbook, monkeypatch):
    from woodchopping.ui import handicap_ui

    path, _ = result_workbook
    monkeypatch.setattr(handicap_ui, "RESULTS_FILE", str(path))
    answers = iter(["2", "25"])
    monkeypatch.setattr("builtins.input", lambda _prompt: next(answers))
    round_object = {"competitors": ["Synthetic Athlete"], "round_name": "=1+1"}
    handicap_ui.append_results_to_excel(
        wood_selection={"event": "UH", "species": "Aspen", "size_mm": 300, "quality": 7},
        round_object=round_object,
        tournament_state={"event_name": "Synthetic"},
    )
    row = _rows(path)[0]
    assert row["CompetitorID"] == ("=1+1", "s")
    assert row["Time (seconds)"] == (25, "n")
    assert row["Round"] == ("=1+1", "s")
    assert round_object["actual_results"] == {"Synthetic Athlete": 25.0}
